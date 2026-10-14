"""Revisioned, permission-filtered lexical source index for class uploads."""

from __future__ import annotations

import asyncio
import difflib
import hashlib
import json
import logging
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Protocol
from uuid import uuid4

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pypdf import PdfReader

from app.api.errors import ApiError
from app.services.class_access import (
    can_view_assessment_key,
    can_view_unreleased,
    class_access,
    is_class_manager,
    is_effectively_published,
    is_private_assessment,
)
from app.services.database import SQLiteConnection, SQLiteDatabase, json_value

logger = logging.getLogger("screen_tutor.source_index")
MAX_SEARCH_QUERY = 500
MAX_SEARCH_TERMS = 16
MAX_SEARCH_LIMIT = 20
MAX_CHUNK_TOKENS = 600
TARGET_CHUNK_TOKENS = 450
MAX_BACKFILL_PER_START = 50
EMBEDDING_INTERFACE_VERSION = "source-embedding-v1"
_WORD = re.compile(r"\w+|[^\w\s]", re.UNICODE)
_MARKER = re.compile(r"^\[(Page|Slide)\s+(\d+)\]\s*$", re.IGNORECASE)
_HEADING = re.compile(r"^#{1,6}\s+(.+?)\s*#*\s*$")


@dataclass(frozen=True)
class SourceBlock:
    """A source excerpt with its original structural and provenance metadata."""

    title: str
    heading_path: tuple[str, ...]
    text: str
    locator: dict[str, Any]
    inferred: bool
    extraction_method: str
    provenance: str = "source"


class EmbeddingProvider(Protocol):
    """Optional provider-neutral embedding seam; no provider is configured in Gate 1."""

    version: str

    async def embed(self, text: str) -> list[float]: ...


def sha256_text(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def file_sha256(path: Path | None) -> str | None:
    if path is None or not path.is_file():
        return None
    digest = hashlib.sha256()
    with path.open("rb") as source:
        while chunk := source.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def _split_sections(text: str) -> list[tuple[tuple[str, ...], str, str]]:
    """Split Markdown/plain text into heading-backed blocks."""
    stack: list[str] = []
    current: list[str] = []
    blocks: list[tuple[tuple[str, ...], str, str]] = []
    current_path: tuple[str, ...] = ("Document",)
    for line in text.splitlines():
        match = _HEADING.match(line.strip())
        if match:
            body = "\n".join(current).strip()
            if body:
                blocks.append((current_path, current_path[-1], body))
            level = len(line) - len(line.lstrip("#"))
            heading = match.group(1).strip()
            stack = stack[: max(0, level - 1)]
            stack.append(heading)
            current_path = tuple(stack)
            current = []
        else:
            current.append(line)
    body = "\n".join(current).strip()
    if body:
        blocks.append((current_path, current_path[-1], body))
    if not blocks and text.strip():
        blocks.append((("Document",), "Document", text.strip()))
    return blocks


def _split_marked_pages(text: str) -> list[SourceBlock]:
    blocks: list[SourceBlock] = []
    current_page: int | None = None
    current: list[str] = []

    def flush() -> None:
        body = "\n".join(current).strip()
        if not body:
            return
        page = current_page
        lines = [line.strip() for line in body.splitlines() if line.strip()]
        title = lines[0][:120] if lines else f"Page {page or len(blocks) + 1}"
        if len(lines) > 1 and len(title) < 100:
            # Use the first line as a title only when it looks like a heading.
            title = (
                title
                if title.isupper() or len(title.split()) <= 9
                else f"Page {page or len(blocks) + 1}"
            )
        blocks.append(
            SourceBlock(
                title=title,
                heading_path=("Document", title),
                text=body,
                locator={"page": page} if page is not None else {},
                inferred=page is None,
                extraction_method="pdf_text",
            )
        )

    for line in text.splitlines():
        match = _MARKER.match(line.strip())
        if match and match.group(1).casefold() == "page":
            flush()
            current = []
            current_page = int(match.group(2))
        else:
            current.append(line)
    flush()
    return blocks


def _pdf_blocks(path: Path, extracted_text: str) -> list[SourceBlock]:
    try:
        pages = PdfReader(str(path)).pages
        native = [page.extract_text() or "" for page in pages]
    except Exception:
        native = []
    vision_blocks = _split_marked_pages(extracted_text)
    if native:
        vision_by_page = {
            int(block.locator["page"]): block
            for block in vision_blocks
            if block.locator.get("page") is not None
        }
        result: list[SourceBlock] = []
        heading_stack: list[str] = []
        for number, native_text in enumerate(native, 1):
            native_text = native_text.strip()
            vision = vision_by_page.get(number)
            # Prefer native page text when it has substance. For mixed/scanned
            # PDFs, use the corresponding extracted vision page where the native
            # layer is empty or only contains a short header/footer.
            chosen = native_text if len(native_text) >= 80 or vision is None else vision.text
            if not chosen:
                continue
            is_native = chosen == native_text
            method = "pdf_text" if is_native else "gemini_pdf_vision"
            if is_native:
                result.extend(_native_pdf_page_blocks(number, chosen, heading_stack))
            else:
                title = next(
                    (line.strip() for line in chosen.splitlines() if line.strip()),
                    f"Page {number}",
                )
                result.append(
                    SourceBlock(
                        title=title[:120],
                        heading_path=("Document", title[:120]),
                        text=chosen,
                        locator={"page": number},
                        inferred=True,
                        extraction_method=method,
                    )
                )
        if result:
            return result
    if vision_blocks:
        return [
            SourceBlock(
                **{**block.__dict__, "extraction_method": "gemini_pdf_vision", "inferred": True}
            )
            for block in vision_blocks
        ]
    return _plain_blocks(extracted_text, method="gemini_pdf_vision", inferred=True)


def _native_pdf_page_blocks(
    page_number: int, text: str, heading_stack: list[str]
) -> list[SourceBlock]:
    """Split native page text at explicit Markdown headings, retaining page locators."""
    blocks: list[SourceBlock] = []
    current: list[str] = []
    saw_heading = False

    def flush() -> None:
        body = "\n".join(current).strip()
        if not body:
            return
        path = tuple(heading_stack) or ("Document", f"Page {page_number}")
        blocks.append(
            SourceBlock(
                title=path[-1][:120],
                heading_path=path,
                text=body,
                locator={"page": page_number},
                inferred=False,
                extraction_method="pdf_text",
            )
        )

    for line in text.splitlines():
        match = _HEADING.match(line.strip())
        if match:
            flush()
            saw_heading = True
            level = len(line) - len(line.lstrip("#"))
            heading_stack[:] = heading_stack[: max(0, level - 1)]
            heading_stack.append(match.group(1).strip())
            current = []
        else:
            current.append(line)
    flush()

    # Preserve native page text with no explicit hierarchy as one page block.
    # A known heading from the prior page remains its context when the PDF text
    # stream continues a section across a page boundary.
    if not blocks and text.strip() and not saw_heading:
        path = tuple(heading_stack) or ("Document", f"Page {page_number}")
        blocks.append(
            SourceBlock(
                title=path[-1][:120],
                heading_path=path,
                text=text.strip(),
                locator={"page": page_number},
                inferred=False,
                extraction_method="pdf_text",
            )
        )
    return blocks


def _docx_blocks(path: Path) -> list[SourceBlock]:
    document = Document(str(path))
    headings: list[str] = []
    sections: dict[tuple[str, ...], list[tuple[str, int]]] = {}
    order: list[tuple[str, ...]] = []
    table_number = 0
    block_order = 0
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style = paragraph.style.name if paragraph.style else ""
            match = re.match(r"Heading\s+(\d+)", style, re.IGNORECASE)
            if match:
                level = min(int(match.group(1)), 6)
                headings = headings[: level - 1]
                headings.append(text[:200])
                path = tuple(headings)
                if path not in sections:
                    sections[path] = []
                    order.append(path)
                continue
            path = tuple(headings) or ("Document",)
            if path not in sections:
                sections[path] = []
                order.append(path)
            sections[path].append((text, block_order))
            block_order += 1
        elif child.tag.endswith("}tbl"):
            table_number += 1
            table = Table(child, document)
            path = tuple(headings) or ("Document",)
            if path not in sections:
                sections[path] = []
                order.append(path)
            for row_number, row in enumerate(table.rows, 1):
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(cells):
                    sections[path].append(
                        (
                            " | ".join(cells) + f" [table {table_number}, row {row_number}]",
                            block_order,
                        )
                    )
                    block_order += 1
    return [
        SourceBlock(
            title=path[-1],
            heading_path=path,
            text="\n\n".join(text for text, _ in sections[path]),
            locator={
                "block_order_start": sections[path][0][1],
                "block_order_end": sections[path][-1][1],
            },
            inferred=path == ("Document",),
            extraction_method="docx_structure",
        )
        for path in order
        if sections[path]
    ]


def _pptx_blocks(path: Path, extracted_text: str) -> list[SourceBlock]:
    presentation = Presentation(str(path))
    blocks: list[SourceBlock] = []
    for number, slide in enumerate(presentation.slides, 1):
        shapes = [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False)]
        title_shape = next((shape for shape in shapes if getattr(shape, "is_title", False)), None)
        lines = [shape.text.strip() for shape in shapes if shape.text.strip()]
        if not lines:
            continue
        title = (
            title_shape.text.strip() if title_shape and title_shape.text.strip() else lines[0]
        )[:160]
        blocks.append(
            SourceBlock(
                title=title,
                heading_path=("Slides", title),
                text="\n".join(lines),
                locator={"slide": number},
                inferred=False,
                extraction_method="pptx_shapes",
            )
        )
    marker = "[Slide image descriptions]"
    if marker in extracted_text:
        description = extracted_text.split(marker, 1)[1].strip()
        if description:
            description_blocks = _split_pptx_image_descriptions(description)
            blocks.extend(description_blocks)
    return blocks or _plain_blocks(extracted_text, method="pptx_extracted", inferred=True)


def _split_pptx_image_descriptions(text: str) -> list[SourceBlock]:
    """Bind generated slide descriptions to slide labels supplied to the model."""
    slide_marker = re.compile(r"^\[Slide\s+(\d+)\]\s*$", re.IGNORECASE)
    descriptions: list[tuple[int | None, list[str]]] = []
    current_slide: int | None = None
    current: list[str] = []

    def flush() -> None:
        body = "\n".join(current).strip()
        if body:
            descriptions.append((current_slide, [body]))

    for line in text.splitlines():
        match = slide_marker.match(line.strip())
        if match:
            flush()
            current_slide = int(match.group(1))
            current = []
        else:
            current.append(line)
    flush()

    # Merge repeated labels for slides containing multiple images while keeping
    # unlabelled provider output explicitly unlocated.
    grouped: dict[int | None, list[str]] = {}
    order: list[int | None] = []
    for slide, bodies in descriptions:
        if slide not in grouped:
            grouped[slide] = []
            order.append(slide)
        grouped[slide].extend(bodies)

    blocks: list[SourceBlock] = []
    for slide in order:
        title = (
            f"Slide {slide} image description"
            if slide is not None
            else "Unlocated image description"
        )
        blocks.append(
            SourceBlock(
                title=title,
                heading_path=("Generated visual descriptions", title),
                text="\n\n".join(grouped[slide]),
                locator={"slide": slide} if slide is not None else {},
                inferred=True,
                extraction_method="gemini_vision_description",
                provenance="generated_description",
            )
        )
    return blocks


def _plain_blocks(text: str, *, method: str, inferred: bool = False) -> list[SourceBlock]:
    return [
        SourceBlock(
            title=title,
            heading_path=path,
            text=body,
            locator={},
            inferred=inferred,
            extraction_method=method,
        )
        for path, title, body in _split_sections(text)
    ]


def make_source_blocks(
    *,
    path: Path | None,
    kind: str,
    source_type: str,
    mime_type: str | None,
    extracted_text: str,
    transcript_segments: list[dict[str, Any]] | None = None,
    extraction_method: str | None = None,
) -> list[SourceBlock]:
    """Extract locators and hierarchy without making another provider request."""
    if transcript_segments:
        blocks = []
        current_text: list[str] = []
        current_tokens = 0
        first_start: float | None = None
        last_end: float | None = None
        first_index = 0
        last_index = 0

        def flush_segment_group() -> None:
            nonlocal current_text, current_tokens, first_start, last_end, first_index, last_index
            if not current_text:
                return
            blocks.append(
                SourceBlock(
                    title=f"Transcript segments {first_index}–{last_index}",
                    heading_path=("Transcript", f"Part {len(blocks) + 1}"),
                    text=" ".join(current_text),
                    locator={
                        "timestamp_start_seconds": first_start or 0.0,
                        "timestamp_end_seconds": last_end or first_start or 0.0,
                    },
                    inferred=False,
                    extraction_method="lecture_transcript",
                )
            )
            current_text, current_tokens = [], 0
            first_start, last_end = None, None
            first_index = last_index = 0

        for index, segment in enumerate(transcript_segments, 1):
            text = str(segment.get("text") or "").strip()
            if not text:
                continue
            start = float(segment.get("start_seconds", 0))
            end = float(segment.get("end_seconds", start))
            token_count = len(_WORD.findall(text))
            if current_text and current_tokens + token_count > MAX_CHUNK_TOKENS:
                flush_segment_group()
            if not current_text:
                first_start = start
                first_index = index
            current_text.append(text)
            current_tokens += token_count
            last_end = end
            last_index = index
            if current_tokens >= TARGET_CHUNK_TOKENS:
                flush_segment_group()
        flush_segment_group()
        return blocks
    if path is None or source_type == "pasted_text":
        return _plain_blocks(
            extracted_text,
            method="pasted_text" if source_type == "pasted_text" else "text",
            inferred=False,
        )
    suffix = path.suffix.casefold()
    try:
        if suffix == ".pdf":
            blocks = _pdf_blocks(path, extracted_text)
        elif suffix == ".pptx":
            blocks = _pptx_blocks(path, extracted_text)
        elif suffix == ".docx":
            blocks = _docx_blocks(path)
        elif suffix in {".txt", ".md"}:
            blocks = _plain_blocks(extracted_text, method="text_file", inferred=False)
        else:
            blocks = _plain_blocks(
                extracted_text,
                method=extraction_method
                or (
                    "gemini_vision" if suffix in {".png", ".jpg", ".jpeg"} else "gemini_extraction"
                ),
                inferred=True,
            )
        return blocks or _plain_blocks(
            extracted_text, method=extraction_method or "extracted_text", inferred=True
        )
    except Exception as exc:
        logger.warning("source_structure_fallback category=%s", type(exc).__name__)
        return _plain_blocks(
            extracted_text, method=extraction_method or "extracted_text", inferred=True
        )


def _split_block(block: SourceBlock) -> list[tuple[str, int]]:
    text = block.text.strip()
    matches = list(_WORD.finditer(text))
    if len(matches) <= MAX_CHUNK_TOKENS:
        return [(text, len(matches))] if text else []
    result: list[tuple[str, int]] = []
    char_start = 0
    token_start = 0
    while token_start < len(matches):
        token_end = min(token_start + TARGET_CHUNK_TOKENS, len(matches))
        char_end = matches[token_end - 1].end()
        chunk = text[char_start:char_end]
        if chunk.strip():
            result.append((chunk, token_end - token_start))
        char_start = char_end
        token_start = token_end
    if char_start < len(text):
        trailing = text[char_start:]
        if result and len(_WORD.findall(trailing)) == 0:
            previous, count = result[-1]
            result[-1] = (previous + trailing, count)
        elif trailing.strip():
            result.append((trailing, len(_WORD.findall(trailing))))
    return result


def _normalize(value: str) -> str:
    return " ".join(re.findall(r"\w+", value.casefold(), re.UNICODE))


def _similarity(left: str, right: str) -> float:
    left_n, right_n = _normalize(left), _normalize(right)
    if not left_n or not right_n:
        return 0.0
    if left_n == right_n:
        return 1.0
    return difflib.SequenceMatcher(None, left_n, right_n, autojunk=False).ratio()


def _locator_kind(locator: dict[str, Any]) -> str:
    return next((key for key in ("page", "slide", "timestamp_start_seconds") if key in locator), "")


def _old_section_content(rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for row in rows:
        entry = result.setdefault(
            row["section_id"],
            {
                "title": row["title"],
                "path": json_value(row["heading_path"], []),
                "locator": json_value(row["locator"], {}),
                "texts": [],
            },
        )
        if row.get("text"):
            entry["texts"].append(row["text"])
    for entry in result.values():
        entry["text"] = "\n\n".join(entry.pop("texts"))
    return result


def _match_sections(
    sections: list[dict[str, Any]], old_sections: dict[str, dict[str, Any]]
) -> dict[tuple[str, ...], str]:
    unused = set(old_sections)
    mapped: dict[tuple[str, ...], str] = {}
    for section in sections:
        path = tuple(section["path"])
        best_id = None
        best_score = 0.0
        new_text = "\n\n".join(section["texts"])
        for section_id in sorted(unused):
            old = old_sections[section_id]
            old_path = tuple(old["path"] or [])
            path_score = difflib.SequenceMatcher(
                None, _normalize(" / ".join(path)), _normalize(" / ".join(old_path)), autojunk=False
            ).ratio()
            title_score = _similarity(section["title"], old["title"])
            text_score = _similarity(new_text[:5000], old["text"][:5000])
            locator_bonus = (
                0.02
                if _locator_kind(section["locator"]) == _locator_kind(old["locator"])
                and _locator_kind(section["locator"])
                else 0.0
            )
            score = 0.40 * path_score + 0.25 * title_score + 0.35 * text_score + locator_bonus
            if score > best_score:
                best_id, best_score = section_id, score
        if best_id is not None and best_score >= 0.58:
            mapped[path] = best_id
            unused.remove(best_id)
    return mapped


async def persist_revision(
    connection: SQLiteConnection,
    *,
    note_id: str,
    processing_revision: int,
    source_hash: str | None,
    extracted_text: str,
    extraction_method: str,
    blocks: list[SourceBlock],
    mark_note_ready: bool = False,
    upload_result: dict[str, Any] | None = None,
) -> str | None:
    """Persist all revision data atomically, discarding results for stale source jobs."""
    note = await connection.fetchrow(
        "SELECT id,status,processing_revision,source_type FROM course_notes WHERE id=?1", note_id
    )
    expected_status = "processing" if mark_note_ready else "ready"
    if (
        note is None
        or note["status"] != expected_status
        or int(note["processing_revision"]) != processing_revision
    ):
        await connection.execute(
            "UPDATE source_index_jobs SET "
            "status='superseded',last_error_code='revision_superseded',"
            "completed_at=strftime('%Y-%m-%dT%H:%M:%fZ','now'),"
            "updated_at=strftime('%Y-%m-%dT%H:%M:%fZ','now') "
            "WHERE note_id=?1 AND processing_revision=?2 AND status IN ('running','queued')",
            note_id,
            processing_revision,
        )
        return None
    job = await connection.fetchrow(
        "SELECT id FROM source_index_jobs WHERE note_id=?1 AND "
        "processing_revision=?2 AND status='running'",
        note_id,
        processing_revision,
    )
    if job is None:
        return None

    previous = await connection.fetchrow(
        "SELECT r.id,r.id AS revision_id,r.source_sha256,r.extracted_sha256 "
        "FROM source_index_heads h "
        "JOIN source_revisions r ON r.id=h.revision_id WHERE h.note_id=?1",
        note_id,
    )
    extracted_hash = sha256_text(extracted_text)
    if (
        previous is not None
        and previous["source_sha256"] == source_hash
        and previous["extracted_sha256"] == extracted_hash
    ):
        await connection.execute(
            "UPDATE source_index_heads SET processing_revision=?2,"
            "updated_at=strftime('%Y-%m-%dT%H:%M:%fZ','now') WHERE note_id=?1",
            note_id,
            processing_revision,
        )
        if mark_note_ready:
            result = upload_result or {}
            await connection.execute(
                "UPDATE course_notes SET "
                "extracted_text=?3,summary=?4,key_topics=?5,transcript_segments=?6,"
                "duration_seconds=?7,syllabus_outline=?8,source_index_status='ready',"
                "status='ready',processing_progress=100,failure_code=NULL,"
                "processed_at=strftime('%Y-%m-%dT%H:%M:%fZ','now'),"
                "updated_at=strftime('%Y-%m-%dT%H:%M:%fZ','now') "
                "WHERE id=?1 AND status='processing' AND processing_revision=?2",
                note_id,
                processing_revision,
                result.get("extracted_text", extracted_text),
                result.get("summary"),
                json.dumps(result.get("key_topics", [])),
                json.dumps(result.get("transcript_segments", [])),
                result.get("duration_seconds"),
                json.dumps(result["syllabus_outline"])
                if result.get("syllabus_outline") is not None
                else None,
            )
        else:
            await connection.execute(
                "UPDATE course_notes SET source_index_status='ready',updated_at="
                "strftime('%Y-%m-%dT%H:%M:%fZ','now') "
                "WHERE id=?1 AND status='ready' AND processing_revision=?2",
                note_id,
                processing_revision,
            )
        await connection.execute(
            "UPDATE source_index_jobs SET status='ready',last_error_code=NULL,completed_at="
            "strftime('%Y-%m-%dT%H:%M:%fZ','now'),updated_at=strftime('%Y-%m-%dT%H:%M:%fZ','now') "
            "WHERE note_id=?1 AND processing_revision=?2 AND status='running'",
            note_id,
            processing_revision,
        )
        return previous["id"]

    old_sections: dict[str, dict[str, Any]] = {}
    old_chunks: dict[str, list[dict[str, Any]]] = {}
    if previous is not None:
        old_rows = await connection.fetch(
            "SELECT sr.section_id,sr.title,sr.heading_path,sr.locator,cr.text "
            "FROM source_section_revisions sr LEFT JOIN source_chunk_revisions cr "
            "ON cr.revision_id=sr.revision_id AND cr.section_id=sr.section_id "
            "WHERE sr.revision_id=?1 ORDER BY sr.section_order,cr.chunk_order",
            previous["revision_id"],
        )
        old_sections = _old_section_content(old_rows)
        chunk_rows = await connection.fetch(
            "SELECT section_id,chunk_id,text,locator FROM source_chunk_revisions "
            "WHERE revision_id=?1 ORDER BY section_id,chunk_order",
            previous["revision_id"],
        )
        for chunk in chunk_rows:
            old_chunks.setdefault(chunk["section_id"], []).append(chunk)

    _, normalized = normalize_blocks(blocks)
    matched = _match_sections(normalized, old_sections)
    revision_number = int(
        await connection.fetchval(
            "SELECT COALESCE(MAX(revision_number),0)+1 FROM source_revisions WHERE note_id=?1",
            note_id,
        )
    )
    revision_id = str(uuid4())
    await connection.execute(
        "INSERT INTO source_revisions(id,note_id,revision_number,processing_revision,source_sha256,"
        "extracted_sha256,extraction_method) VALUES (?1,?2,?3,?4,?5,?6,?7)",
        revision_id,
        note_id,
        revision_number,
        processing_revision,
        source_hash,
        extracted_hash,
        extraction_method,
    )

    path_to_id: dict[tuple[str, ...], str] = {}
    for section in normalized:
        path = tuple(section["path"])
        existing_id = matched.get(path)
        if existing_id:
            section_id = existing_id
        else:
            section_id = str(uuid4())
            await connection.execute(
                "INSERT INTO source_sections(id,note_id) VALUES (?1,?2)", section_id, note_id
            )
        path_to_id[path] = section_id
    section_order = {tuple(section["path"]): index for index, section in enumerate(normalized)}
    for section in normalized:
        path = tuple(section["path"])
        section_id = path_to_id[path]
        parent_id = path_to_id.get(path[:-1]) if len(path) > 1 else None
        await connection.execute(
            "INSERT INTO "
            "source_section_revisions(revision_id,section_id,parent_section_id,section_order,"
            "title,heading_path,locator,inferred) VALUES (?1,?2,?3,?4,?5,?6,?7,?8)",
            revision_id,
            section_id,
            parent_id,
            section_order[path],
            section["title"],
            json.dumps(list(path), ensure_ascii=False),
            json.dumps(section["locator"], ensure_ascii=False),
            int(section["inferred"]),
        )

    global_chunk_order = 0
    for section in normalized:
        section_id = path_to_id[tuple(section["path"])]
        prior = old_chunks.get(section_id, [])
        used: set[str] = set()
        chunks = section["chunks"]
        for chunk in chunks:
            chosen = None
            chosen_score = 0.0
            for old in prior:
                if old["chunk_id"] in used:
                    continue
                score = _similarity(chunk["text"], old["text"])
                if _locator_kind(chunk["locator"]) == _locator_kind(
                    json_value(old["locator"], {})
                ) and _locator_kind(chunk["locator"]):
                    score += 0.02
                if score > chosen_score:
                    chosen, chosen_score = old, score
            if chosen is not None and chosen_score >= 0.62:
                logical_chunk_id = chosen["chunk_id"]
                used.add(logical_chunk_id)
            else:
                logical_chunk_id = str(uuid4())
                await connection.execute(
                    "INSERT INTO source_chunks(id,section_id) VALUES (?1,?2)",
                    logical_chunk_id,
                    section_id,
                )
            chunk_revision_id = str(uuid4())
            await connection.execute(
                "INSERT INTO "
                "source_chunk_revisions(id,revision_id,section_id,chunk_id,chunk_order,locator,"
                "text,content_sha256,token_count,extraction_method,provenance) "
                "VALUES (?1,?2,?3,?4,?5,?6,?7,?8,?9,?10,?11)",
                chunk_revision_id,
                revision_id,
                section_id,
                logical_chunk_id,
                global_chunk_order,
                json.dumps(chunk["locator"], ensure_ascii=False),
                chunk["text"],
                sha256_text(chunk["text"]),
                chunk["token_count"],
                chunk["extraction_method"],
                chunk["provenance"],
            )
            global_chunk_order += 1

    await connection.execute(
        "INSERT INTO source_index_heads(note_id,revision_id,processing_revision) "
        "VALUES (?1,?2,?3) "
        "ON CONFLICT(note_id) DO UPDATE SET revision_id=excluded.revision_id,"
        "processing_revision=excluded.processing_revision,"
        "updated_at=strftime('%Y-%m-%dT%H:%M:%fZ','now')",
        note_id,
        revision_id,
        processing_revision,
    )
    if mark_note_ready:
        result = upload_result or {}
        await connection.execute(
            "UPDATE course_notes SET "
            "extracted_text=?3,summary=?4,key_topics=?5,transcript_segments=?6,"
            "duration_seconds=?7,syllabus_outline=?8,source_index_status='ready',status='ready',"
            "processing_progress=100,failure_code=NULL,"
            "processed_at=strftime('%Y-%m-%dT%H:%M:%fZ','now'),"
            "updated_at=strftime('%Y-%m-%dT%H:%M:%fZ','now') "
            "WHERE id=?1 AND status='processing' AND processing_revision=?2",
            note_id,
            processing_revision,
            result.get("extracted_text", extracted_text),
            result.get("summary"),
            json.dumps(result.get("key_topics", [])),
            json.dumps(result.get("transcript_segments", [])),
            result.get("duration_seconds"),
            json.dumps(result["syllabus_outline"])
            if result.get("syllabus_outline") is not None
            else None,
        )
    else:
        await connection.execute(
            "UPDATE course_notes SET "
            "source_index_status='ready',updated_at=strftime('%Y-%m-%dT%H:%M:%fZ','now') "
            "WHERE id=?1 AND status='ready' AND processing_revision=?2",
            note_id,
            processing_revision,
        )
    await connection.execute(
        "UPDATE source_index_jobs SET status='ready',last_error_code=NULL,completed_at="
        "strftime('%Y-%m-%dT%H:%M:%fZ','now'),updated_at=strftime('%Y-%m-%dT%H:%M:%fZ','now') "
        "WHERE note_id=?1 AND processing_revision=?2 AND status='running'",
        note_id,
        processing_revision,
    )
    return revision_id


def normalize_blocks(
    blocks: list[SourceBlock],
) -> tuple[dict[tuple[str, ...], list[dict]], list[dict[str, Any]]]:
    section_meta: dict[tuple[str, ...], dict[str, Any]] = {}
    sections: list[dict[str, Any]] = []
    for block in blocks:
        path = block.heading_path or (block.title or "Document",)
        for depth in range(1, len(path) + 1):
            prefix = path[:depth]
            if prefix not in section_meta:
                section = {
                    "path": list(prefix),
                    "title": prefix[-1][:200],
                    "locator": block.locator if depth == len(path) else {},
                    "inferred": block.inferred if depth == len(path) else False,
                    "texts": [],
                    "chunks": [],
                }
                section_meta[prefix] = section
                sections.append(section)
        section_meta[path]["title"] = block.title[:200] or path[-1]
        section_meta[path]["locator"] = block.locator
        section_meta[path]["inferred"] = block.inferred
        for text, token_count in _split_block(block):
            section_meta[path]["texts"].append(text)
            section_meta[path]["chunks"].append(
                {
                    "text": text,
                    "token_count": token_count,
                    "locator": block.locator,
                    "extraction_method": block.extraction_method,
                    "provenance": block.provenance,
                }
            )
    # Insertion order follows source order and parents are inserted before their first child.
    return {}, sections


async def begin_index_job(database: SQLiteDatabase, note_id: str, processing_revision: int) -> bool:
    async with database.acquire() as connection:
        async with connection.transaction():
            note = await connection.fetchrow(
                "SELECT id FROM course_notes WHERE id=?1 AND status='processing' "
                "AND processing_revision=?2",
                note_id,
                processing_revision,
            )
            if note is None:
                return False
            job = await connection.fetchrow(
                "INSERT INTO source_index_jobs(id,note_id,processing_revision,status,attempts) "
                "VALUES (?1,?2,?3,'running',1) ON "
                "CONFLICT(note_id,processing_revision) DO UPDATE SET "
                "status='running',attempts=source_index_jobs.attempts+1,last_error_code=NULL,"
                "started_at=strftime('%Y-%m-%dT%H:%M:%fZ','now'),completed_at=NULL,"
                "updated_at=strftime('%Y-%m-%dT%H:%M:%fZ','now') "
                "WHERE source_index_jobs.status IN ('queued','failed') "
                "AND source_index_jobs.attempts < source_index_jobs.max_attempts RETURNING id",
                uuid4(),
                note_id,
                processing_revision,
            )
            if job is None:
                return False
            await connection.execute(
                "UPDATE course_notes SET source_index_status='processing' "
                "WHERE id=?1 AND status='processing' AND processing_revision=?2",
                note_id,
                processing_revision,
            )
            return True


async def enqueue_backfill_jobs(
    database: SQLiteDatabase, *, limit: int = MAX_BACKFILL_PER_START
) -> int:
    """Queue a bounded page of ready legacy uploads lacking a source index."""
    rows = await database.fetch(
        "SELECT n.id,n.processing_revision FROM course_notes n "
        "LEFT JOIN source_index_heads h ON h.note_id=n.id "
        "WHERE n.status='ready' AND (h.note_id IS NULL OR "
        "h.processing_revision<>n.processing_revision) AND NOT EXISTS "
        "(SELECT 1 FROM source_index_jobs j WHERE j.note_id=n.id AND "
        "j.processing_revision=n.processing_revision) "
        "AND NOT EXISTS (SELECT 1 FROM uploaded_assessments a WHERE a.note_id=n.id AND "
        "(a.has_student_responses=1 OR a.score_percent IS NOT NULL OR "
        "a.grading_result IS NOT NULL)) "
        "ORDER BY n.created_at,n.id LIMIT ?1",
        limit,
    )
    if not rows:
        return 0
    async with database.acquire() as connection:
        async with connection.transaction():
            await connection.executemany(
                "INSERT OR IGNORE INTO source_index_jobs(id,note_id,processing_revision,status) "
                "VALUES (?1,?2,?3,'queued')",
                [(uuid4(), row["id"], row["processing_revision"]) for row in rows],
            )
            await connection.executemany(
                "UPDATE course_notes SET source_index_status='processing' WHERE id=?1 "
                "AND status='ready'",
                [(row["id"],) for row in rows],
            )
    return len(rows)


async def recover_source_index_jobs(database: SQLiteDatabase) -> None:
    """Bound recovery to one indexed batch; stale running jobs become retryable."""
    async with database.acquire() as connection:
        async with connection.transaction():
            await connection.execute(
                "UPDATE source_index_jobs SET "
                "status='failed',last_error_code='upload_processing_interrupted',"
                "completed_at=strftime('%Y-%m-%dT%H:%M:%fZ','now'),"
                "updated_at=strftime('%Y-%m-%dT%H:%M:%fZ','now') WHERE status IN "
                "('queued','running') "
                "AND EXISTS (SELECT 1 FROM course_notes n WHERE "
                "n.id=source_index_jobs.note_id AND n.status='failed')"
            )
            await connection.execute(
                "UPDATE source_index_jobs SET status='queued',last_error_code='retry_recovered',"
                "completed_at=NULL,updated_at=strftime('%Y-%m-%dT%H:%M:%fZ','now') "
                "WHERE status='failed' AND attempts < max_attempts AND last_error_code IN "
                "('index_failed','worker_interrupted')"
            )
            await connection.execute(
                "UPDATE source_index_jobs SET status=CASE WHEN attempts < "
                "max_attempts THEN 'queued' ELSE 'failed' END,"
                "last_error_code='worker_interrupted',"
                "updated_at=strftime('%Y-%m-%dT%H:%M:%fZ','now') "
                "WHERE status='running'"
            )
            await connection.execute(
                "UPDATE course_notes SET source_index_status=CASE WHEN EXISTS ("
                "SELECT 1 FROM source_index_jobs j WHERE j.note_id=course_notes.id "
                "AND j.status='queued') "
                "THEN 'processing' ELSE 'failed' END WHERE source_index_status='processing'"
            )
    await enqueue_backfill_jobs(database)


async def process_queued_jobs(
    application: Any, *, limit: int = MAX_BACKFILL_PER_START, note_id: str | None = None
) -> int:
    """Process a bounded batch of no-provider indexing work after startup or search."""
    database: SQLiteDatabase | None = getattr(application.state, "database", None)
    storage = getattr(application.state, "private_storage", None)
    if database is None:
        return 0
    await database.execute(
        "UPDATE source_index_jobs SET status='superseded',last_error_code='revision_superseded',"
        "completed_at=strftime('%Y-%m-%dT%H:%M:%fZ','now'),"
        "updated_at=strftime('%Y-%m-%dT%H:%M:%fZ','now') "
        "WHERE status='queued' AND EXISTS (SELECT 1 FROM course_notes n WHERE "
        "n.id=source_index_jobs.note_id "
        "AND (n.status<>'ready' OR n.processing_revision<>source_index_jobs.processing_revision))"
    )
    rows = await database.fetch(
        "SELECT j.id AS "
        "job_id,j.note_id,j.processing_revision,n.title,n.kind,n.source_type,n.storage_key,"
        "n.mime_type,n.extracted_text,n.transcript_segments FROM source_index_jobs j "
        "JOIN course_notes n ON n.id=j.note_id WHERE j.status='queued' AND n.status='ready' "
        "AND (?2 IS NULL OR j.note_id=?2) "
        "AND n.processing_revision=j.processing_revision AND NOT EXISTS "
        "(SELECT 1 FROM uploaded_assessments a WHERE a.note_id=n.id AND "
        "(a.has_student_responses=1 OR a.score_percent IS NOT NULL OR "
        "a.grading_result IS NOT NULL)) "
        "ORDER BY j.created_at,j.id LIMIT ?1",
        limit,
        note_id,
    )
    completed = 0
    for row in rows:
        async with database.acquire() as connection:
            async with connection.transaction():
                claimed = await connection.fetchrow(
                    "UPDATE source_index_jobs SET status='running',attempts=attempts+1,"
                    "started_at=strftime('%Y-%m-%dT%H:%M:%fZ','now'),"
                    "updated_at=strftime('%Y-%m-%dT%H:%M:%fZ','now') "
                    "WHERE id=?1 AND status='queued' AND attempts < max_attempts RETURNING id",
                    row["job_id"],
                )
                if claimed is None:
                    continue
                await connection.execute(
                    "UPDATE course_notes SET source_index_status='processing' WHERE id=?1 "
                    "AND status='ready' AND processing_revision=?2",
                    row["note_id"],
                    row["processing_revision"],
                )
        path = None
        if row["storage_key"] and storage is not None:
            try:
                path = storage.path_for(row["storage_key"])
            except ValueError:
                path = None
        try:
            blocks = await asyncio.to_thread(
                make_source_blocks,
                path=path,
                kind=row["kind"],
                source_type=row["source_type"],
                mime_type=row["mime_type"],
                extracted_text=row["extracted_text"] or "",
                transcript_segments=json_value(row["transcript_segments"], []),
            )
            source_hash = await asyncio.to_thread(file_sha256, path)
            async with database.acquire() as connection:
                async with connection.transaction():
                    revision = await persist_revision(
                        connection,
                        note_id=row["note_id"],
                        processing_revision=row["processing_revision"],
                        source_hash=source_hash,
                        extracted_text=row["extracted_text"] or "",
                        extraction_method="backfill",
                        blocks=blocks,
                        mark_note_ready=False,
                    )
            if revision:
                completed += 1
        except Exception as exc:
            logger.error(
                "source_index_backfill_failed note_id=%s category=%s",
                row["note_id"],
                type(exc).__name__,
            )
            await database.execute(
                "UPDATE source_index_jobs SET status=CASE WHEN attempts < "
                "max_attempts THEN 'queued' ELSE 'failed' END,"
                "last_error_code='index_failed',completed_at=CASE WHEN attempts < "
                "max_attempts THEN NULL "
                "ELSE strftime('%Y-%m-%dT%H:%M:%fZ','now') "
                "END,updated_at=strftime('%Y-%m-%dT%H:%M:%fZ','now') "
                "WHERE id=?1 AND status='running'",
                row["job_id"],
            )
            await database.execute(
                "UPDATE course_notes SET source_index_status=CASE WHEN EXISTS (SELECT "
                "1 FROM source_index_jobs j "
                "WHERE j.note_id=?1 AND j.processing_revision=?2 AND "
                "j.status='queued') THEN 'processing' ELSE 'failed' END "
                "WHERE id=?1 AND processing_revision=?2",
                row["note_id"],
                row["processing_revision"],
            )
    return completed


async def queue_ready_note_index(
    database: SQLiteDatabase, *, note_id: str, processing_revision: int
) -> None:
    """Queue a newly completed, eligible upload/assessment for local indexing."""
    async with database.acquire() as connection:
        async with connection.transaction():
            note = await connection.fetchrow(
                "SELECT id FROM course_notes WHERE id=?1 AND status='ready' AND "
                "processing_revision=?2",
                note_id,
                processing_revision,
            )
            if note is None:
                return
            await connection.execute(
                "UPDATE course_notes SET source_index_status='processing' WHERE id=?1 "
                "AND status='ready' "
                "AND processing_revision=?2",
                note_id,
                processing_revision,
            )
            await connection.execute(
                "INSERT OR IGNORE INTO source_index_jobs(id,note_id,processing_revision,status) "
                "VALUES (?1,?2,?3,'queued')",
                uuid4(),
                note_id,
                processing_revision,
            )


async def retry_ready_note_index(
    connection: SQLiteConnection, *, note_id: str, processing_revision: int
) -> bool:
    """Reset an exhausted local indexing job without re-running provider extraction."""
    note = await connection.fetchrow(
        "SELECT id FROM course_notes WHERE id=?1 AND status='ready' "
        "AND source_index_status='failed' AND processing_revision=?2",
        note_id,
        processing_revision,
    )
    if note is None:
        return False
    job = await connection.fetchrow(
        "UPDATE source_index_jobs SET status='queued',attempts=0,last_error_code=NULL,"
        "started_at=NULL,completed_at=NULL,updated_at=strftime('%Y-%m-%dT%H:%M:%fZ','now') "
        "WHERE note_id=?1 AND processing_revision=?2 AND status='failed' RETURNING id",
        note_id,
        processing_revision,
    )
    if job is None:
        await connection.execute(
            "INSERT INTO source_index_jobs(id,note_id,processing_revision,status) "
            "VALUES (?1,?2,?3,'queued')",
            uuid4(),
            note_id,
            processing_revision,
        )
    await connection.execute(
        "UPDATE course_notes SET source_index_status='processing' WHERE id=?1 "
        "AND processing_revision=?2",
        note_id,
        processing_revision,
    )
    return True


def fts_query(query: str) -> str:
    terms = re.findall(r"[\w]+", query, re.UNICODE)[:MAX_SEARCH_TERMS]
    if not terms:
        raise ApiError(400, "invalid_query", "Enter a search query containing letters or numbers.")
    return " OR ".join('"' + term.replace('"', '""') + '"' for term in terms)


async def search_authorized_sources(
    database: SQLiteDatabase,
    *,
    class_id: str,
    user_id: str,
    query: str,
    upload_id: str | None,
    limit: int,
) -> dict[str, Any]:
    """Search current indexed chunks after applying live membership/release/key rules."""
    access = await class_access(database, class_id, user_id, required="read")
    if len(query) > MAX_SEARCH_QUERY:
        raise ApiError(400, "invalid_query", "Search query is too long.")
    if not 1 <= limit <= MAX_SEARCH_LIMIT:
        raise ApiError(400, "invalid_limit", "Search limit must be between 1 and 20.")
    match = fts_query(query)
    if upload_id is not None:
        upload = await database.fetchrow(
            "SELECT n.status,n.created_by,n.publication_state,n.release_at,n.resource_scope,"
            "a.id AS assessment_id,a.owner_id AS assessment_owner,a.has_student_responses,"
            "a.score_percent,a.grading_result,a.answer_key_release_state,a.answer_key_release_at "
            "FROM course_notes n LEFT JOIN uploaded_assessments a ON a.note_id=n.id "
            "WHERE n.id=?1 AND n.class_id=?2",
            upload_id,
            class_id,
        )
        if (
            upload is None
            or upload["status"] != "ready"
            or not (
                is_effectively_published(upload["publication_state"], upload["release_at"])
                or can_view_unreleased(access, user_id, upload["created_by"])
            )
        ):
            raise ApiError(404, "not_found", "The requested source was not found.")
        if upload["assessment_id"] is not None:
            assessment = {
                "owner_id": upload["assessment_owner"],
                "has_student_responses": upload["has_student_responses"],
                "score_percent": upload["score_percent"],
                "grading_result": upload["grading_result"],
            }
            if is_private_assessment(assessment) or not can_view_assessment_key(
                {**assessment, **upload}, access, user_id
            ):
                raise ApiError(404, "not_found", "The requested source was not found.")
    rows = await database.fetch(
        "SELECT cr.id AS chunk_revision_id,cr.revision_id,s.id AS section_id,c.id AS chunk_id,"
        "sr.note_id AS material_id,n.title,cr.chunk_order,cr.locator,cr.extraction_method,"
        "cr.provenance,cr.text,sec.title AS section_title,sec.heading_path,sec.parent_section_id,"
        "sec.section_order,sec.locator AS section_locator,sec.inferred AS "
        "section_inferred,cr.content_sha256,"
        "n.created_by,n.publication_state,n.release_at,"
        "n.resource_scope,a.id AS assessment_id,a.owner_id AS "
        "assessment_owner,a.has_student_responses,"
        "a.score_percent,a.grading_result,a.answer_key_release_state,a.answer_key_release_at,"
        "bm25(source_chunks_fts) AS rank_score "
        "FROM source_chunks_fts f JOIN source_chunk_revisions cr ON cr.id=f.chunk_revision_id "
        "JOIN source_revisions sr ON sr.id=cr.revision_id JOIN source_chunks c ON c.id=cr.chunk_id "
        "JOIN source_sections s ON s.id=cr.section_id "
        "JOIN source_section_revisions sec ON sec.revision_id=cr.revision_id "
        "AND sec.section_id=s.id "
        "JOIN course_notes n ON n.id=sr.note_id JOIN source_index_heads h ON "
        "h.note_id=n.id AND h.revision_id=sr.id AND h.processing_revision=n.processing_revision "
        "LEFT JOIN uploaded_assessments a ON a.note_id=n.id "
        "WHERE source_chunks_fts MATCH ?1 AND n.class_id=?2 AND n.status='ready' "
        "AND n.source_index_status='ready' "
        "AND (?5 IS NULL OR n.id=?5) AND (n.publication_state='published' OR "
        "(n.publication_state='scheduled' AND julianday(n.release_at)<=julianday('now')) "
        "OR n.created_by=?3 OR ?4=1) AND (a.id IS NULL OR ("
        "COALESCE(a.has_student_responses,0)=0 AND a.score_percent IS NULL "
        "AND a.grading_result IS NULL "
        "AND ((n.resource_scope='student_contribution' AND (a.owner_id=?3 OR n.created_by=?3 OR "
        "a.answer_key_release_state='published' OR (a.answer_key_release_state='scheduled' AND "
        "julianday(a.answer_key_release_at)<=julianday('now')))) OR "
        "(n.resource_scope<>'student_contribution' AND (?4=1 OR a.owner_id=?3 "
        "OR n.created_by=?3 OR "
        "a.answer_key_release_state='published' OR (a.answer_key_release_state='scheduled' AND "
        "julianday(a.answer_key_release_at)<=julianday('now'))))))) "
        "ORDER BY bm25(source_chunks_fts) ASC,cr.chunk_order ASC LIMIT ?6",
        match,
        class_id,
        user_id,
        int(is_class_manager(access)),
        upload_id,
        limit,
    )
    visible = []
    for row in rows:
        if row["assessment_id"] is not None:
            assessment = {
                "owner_id": row["assessment_owner"],
                "has_student_responses": row["has_student_responses"],
                "score_percent": row["score_percent"],
                "grading_result": row["grading_result"],
            }
            if is_private_assessment(assessment) or not can_view_assessment_key(
                {
                    **assessment,
                    "answer_key_release_state": row["answer_key_release_state"],
                    "answer_key_release_at": row["answer_key_release_at"],
                    "created_by": row["created_by"],
                    "resource_scope": row["resource_scope"],
                },
                access,
                user_id,
            ):
                continue
        if not (
            is_effectively_published(row["publication_state"], row["release_at"])
            or can_view_unreleased(access, user_id, row["created_by"])
        ):
            continue
        snippet = " ".join(row["text"].split())
        if len(snippet) > 600:
            snippet = snippet[:597].rstrip() + "…"
        locator = json_value(row["locator"], {})
        visible.append(
            {
                "material_id": row["material_id"],
                "revision_id": row["revision_id"],
                "section_id": row["section_id"],
                "chunk_id": row["chunk_id"],
                "chunk_revision_id": row["chunk_revision_id"],
                "title": row["title"],
                "order": row["chunk_order"],
                "locator": locator,
                "extraction_method": row["extraction_method"],
                "provenance": row["provenance"],
                "score": round(max(0.0, -float(row["rank_score"])), 8),
                "snippet": snippet,
                "section_title": row["section_title"],
                "heading_path": json_value(row["heading_path"], []),
                "parent_section_id": row["parent_section_id"],
                "section_order": row["section_order"],
                "section_locator": json_value(row["section_locator"], {}),
                "section_inferred": bool(row["section_inferred"]),
                "content_sha256": row["content_sha256"],
            }
        )
    visible = visible[:limit]
    return {"items": visible}


async def fetch_authorized_chunk(
    database: SQLiteDatabase, *, class_id: str, user_id: str, chunk_revision_id: str
) -> dict[str, Any]:
    row = await database.fetchrow(
        "SELECT cr.id AS chunk_revision_id,cr.revision_id,cr.chunk_id,cr.section_id,cr.chunk_order,"
        "cr.locator AS "
        "chunk_locator,cr.text,cr.content_sha256,cr.token_count,cr.extraction_method,cr.provenance,"
        "r.note_id AS "
        "material_id,n.title,n.class_id,n.status,n.source_index_status,n.processing_revision,"
        "h.processing_revision AS indexed_processing_revision,n.created_by,"
        "n.publication_state,n.release_at,"
        "n.resource_scope,a.id AS assessment_id,a.owner_id AS "
        "assessment_owner,a.has_student_responses,"
        "a.score_percent,a.grading_result,a.answer_key_release_state,a.answer_key_release_at,"
        "sec.parent_section_id,sec.section_order,sec.title AS section_title,sec.heading_path,"
        "sec.locator AS section_locator,sec.inferred AS section_inferred "
        "FROM source_chunk_revisions cr JOIN source_revisions r ON r.id=cr.revision_id "
        "JOIN source_sections s ON s.id=cr.section_id JOIN source_section_revisions sec "
        "ON sec.revision_id=cr.revision_id AND sec.section_id=cr.section_id "
        "JOIN course_notes n ON n.id=r.note_id JOIN source_index_heads h ON h.note_id=n.id "
        "LEFT JOIN uploaded_assessments a ON a.note_id=n.id "
        "WHERE cr.id=?1 AND n.class_id=?2",
        chunk_revision_id,
        class_id,
    )
    if row is None:
        raise ApiError(404, "not_found", "The requested source chunk was not found.")
    access = await class_access(database, class_id, user_id, required="read")
    if (
        row["status"] != "ready"
        or row["source_index_status"] != "ready"
        or row["processing_revision"] != row["indexed_processing_revision"]
        or not (
            is_effectively_published(row["publication_state"], row["release_at"])
            or can_view_unreleased(access, user_id, row["created_by"])
        )
    ):
        raise ApiError(404, "not_found", "The requested source chunk was not found.")
    if row["assessment_id"] is not None:
        assessment = {
            "owner_id": row["assessment_owner"],
            "has_student_responses": row["has_student_responses"],
            "score_percent": row["score_percent"],
            "grading_result": row["grading_result"],
        }
        if is_private_assessment(assessment) or not can_view_assessment_key(
            {
                **assessment,
                "answer_key_release_state": row["answer_key_release_state"],
                "answer_key_release_at": row["answer_key_release_at"],
                "created_by": row["created_by"],
                "resource_scope": row["resource_scope"],
            },
            access,
            user_id,
        ):
            raise ApiError(404, "not_found", "The requested source chunk was not found.")
    return {
        "material_id": row["material_id"],
        "revision_id": row["revision_id"],
        "section_id": row["section_id"],
        "chunk_id": row["chunk_id"],
        "chunk_revision_id": row["chunk_revision_id"],
        "title": row["title"],
        "order": row["chunk_order"],
        "locator": json_value(row["chunk_locator"], {}),
        "section_title": row["section_title"],
        "heading_path": json_value(row["heading_path"], []),
        "parent_section_id": row["parent_section_id"],
        "section_order": row["section_order"],
        "section_locator": json_value(row["section_locator"], {}),
        "section_inferred": bool(row["section_inferred"]),
        "text": row["text"],
        "content_sha256": row["content_sha256"],
        "token_count": row["token_count"],
        "extraction_method": row["extraction_method"],
        "provenance": row["provenance"],
    }
