"""Gate 1 source revision, indexing, search, and authorization tests."""

from __future__ import annotations

import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch
from uuid import UUID, uuid4

from docx import Document
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter

from app.api.class_library import _process_upload_job, _reprocess_upload_job
from app.config import Settings
from app.main import create_app
from app.services.auth import SESSION_COOKIE
from app.services.concept_graph import _validate_current_chunk_refs
from app.services.database import create_database
from app.services.source_index import (
    SourceBlock,
    _pdf_blocks,
    _split_block,
    begin_index_job,
    make_source_blocks,
    persist_revision,
)

ORIGIN = "http://localhost:5173"


class FakeContentService:
    def process_text(self, text: str, _kind: str) -> dict:
        return {
            "extracted_text": text,
            "summary": "summary",
            "key_topics": [],
            "transcript_segments": [],
            "duration_seconds": None,
            "syllabus_outline": None,
        }


class SourceIndexApiTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        settings = Settings(
            _env_file=None,
            gemini_api_key="test-key-not-used",
            file_url_signing_secret="test-signing-secret-with-at-least-32-characters",
            sqlite_path=":memory:",
            storage_dir=self.temp.name,
            signup_mode="open",
            lean_tool_enabled=False,
        )
        self.client = TestClient(create_app(settings_override=settings), headers={"origin": ORIGIN})
        self.client.__enter__()
        self.addCleanup(self.client.__exit__, None, None, None)
        self.database = self.client.app.state.database
        self.client.app.state.gemini_content_service = FakeContentService()
        self.tokens: dict[str, str] = {}
        self.users: dict[str, dict] = {}
        self.owner = self.add_user("source_owner")
        response = self.call("post", "/api/classes", "source_owner", json={"name": "Source Index"})
        self.assertEqual(response.status_code, 201, response.text)
        self.class_id = response.json()["id"]

    def add_user(self, username: str) -> dict:
        response = self.client.post(
            "/api/auth/signup", json={"username": username, "password": "source-test-pass"}
        )
        self.assertEqual(response.status_code, 201, response.text)
        self.tokens[username] = response.cookies.get(SESSION_COOKIE)
        self.users[username] = response.json()["user"]
        return self.users[username]

    def call(self, method: str, path: str, username: str, **kwargs):
        return getattr(self.client, method)(
            path, cookies={SESSION_COOKIE: self.tokens[username]}, **kwargs
        )

    def add_reader(self, username: str) -> dict:
        self.add_user(username)
        invite = self.call(
            "post",
            f"/api/classes/{self.class_id}/invitations",
            "source_owner",
            json={"capabilities": ["read"], "max_uses": 1, "expires_in_days": 7},
        )
        self.assertEqual(invite.status_code, 201, invite.text)
        joined = self.call(
            "post",
            "/api/class-invitations/redeem",
            username,
            json={"code": invite.json()["code"]},
        )
        self.assertEqual(joined.status_code, 200, joined.text)
        return self.users[username]

    def seed_processing_note(self, text: str, *, state: str = "published") -> str:
        note_id = str(uuid4())
        self.client.portal.call(
            self.database.execute,
            "INSERT INTO course_notes "
            "(id,class_id,created_by,publication_state,title,kind,source_type,"
            "extracted_text,status,processing_progress) VALUES "
            "(?1,?2,?3,?4,'Source notes','document',"
            "'pasted_text',?5,'processing',0)",
            note_id,
            self.class_id,
            self.owner["id"],
            state,
            text,
        )
        return note_id

    def run_upload_processor(self, note_id: str) -> None:
        with patch("app.api.class_library.create_default_document", new_callable=AsyncMock):
            self.client.portal.call(_process_upload_job, self.client.app, UUID(note_id))

    def test_index_job_claim_rejects_duplicate_running_claim(self) -> None:
        note_id = self.seed_processing_note("Claim once.")
        revision = self.client.portal.call(
            self.database.fetchval,
            "SELECT processing_revision FROM course_notes WHERE id=?1",
            note_id,
        )
        first_claim = self.client.portal.call(begin_index_job, self.database, note_id, revision)
        second_claim = self.client.portal.call(begin_index_job, self.database, note_id, revision)
        self.assertTrue(first_claim)
        self.assertFalse(second_claim)
        job = self.client.portal.call(
            self.database.fetchrow,
            "SELECT status,attempts FROM source_index_jobs WHERE note_id=?1",
            note_id,
        )
        self.assertEqual(job, {"status": "running", "attempts": 1})

    def test_upload_revision_search_exact_fetch_and_stale_revision_protection(self) -> None:
        self.add_reader("source_reader")
        original_text = (
            "# Kernel\nThe kernel manages hardware and operating system services. "
            "It allocates CPU and memory for applications."
        )
        note_id = self.seed_processing_note(original_text)
        self.run_upload_processor(note_id)

        note = self.client.portal.call(
            self.database.fetchrow,
            "SELECT status,source_index_status,processing_revision FROM course_notes WHERE id=?1",
            note_id,
        )
        self.assertEqual((note["status"], note["source_index_status"]), ("ready", "ready"))
        self.assertEqual(
            self.client.portal.call(
                self.database.fetchval,
                "SELECT count(*) FROM source_chunks_fts WHERE note_id=?1",
                note_id,
            ),
            1,
        )

        path = f"/api/classes/{self.class_id}/sources/search"
        searched = self.call("get", path, "source_reader", params={"query": "kernel hardware"})
        self.assertEqual(searched.status_code, 200, searched.text)
        self.assertEqual(searched.headers.get("cache-control"), "private, no-store")
        result = searched.json()["items"][0]
        self.assertEqual(result["material_id"], note_id)
        self.assertIn("chunk_revision_id", result)
        fetched = self.call(
            "get",
            f"/api/classes/{self.class_id}/sources/chunks/{result['chunk_revision_id']}",
            "source_reader",
        )
        self.assertEqual(fetched.status_code, 200, fetched.text)
        self.assertEqual(fetched.json()["text"], original_text.split("\n", 1)[1])

        self.add_user("source_outsider")
        denied = self.call("get", path, "source_outsider", params={"query": "kernel"})
        self.assertEqual(denied.status_code, 404)

        changed_text = (
            "# Kernel\nThe kernel manages hardware and operating system services. "
            "It allocates CPU and memory for applications and handles scheduling."
        )
        self.client.portal.call(
            self.database.execute,
            "UPDATE course_notes SET extracted_text=?2,status='processing' WHERE id=?1",
            note_id,
            changed_text,
        )
        before_reprocess = self.client.portal.call(
            self.database.fetchrow,
            "SELECT status,processing_revision FROM course_notes WHERE id=?1",
            note_id,
        )
        self.assertGreater(before_reprocess["processing_revision"], note["processing_revision"])
        self.assertEqual(before_reprocess["status"], "processing")
        with patch("app.api.class_library.create_default_document", new_callable=AsyncMock):
            self.client.portal.call(
                _reprocess_upload_job,
                self.client.app,
                UUID(note_id),
                True,
                before_reprocess["processing_revision"],
                # The source-change trigger increments processing_revision before the worker
                # claims the job.
            )
        heads = self.client.portal.call(
            self.database.fetch,
            "SELECT r.id AS revision_id,r.processing_revision FROM source_index_heads h "
            "JOIN source_revisions r ON r.id=h.revision_id WHERE h.note_id=?1",
            note_id,
        )
        current = self.client.portal.call(
            self.database.fetchrow,
            "SELECT status,source_index_status,processing_revision,extracted_text "
            "FROM course_notes WHERE id=?1",
            note_id,
        )
        self.assertEqual(
            current["status"],
            "ready",
            {
                "current": current,
                "jobs": self.client.portal.call(
                    self.database.fetch,
                    "SELECT processing_revision,status,last_error_code FROM "
                    "source_index_jobs WHERE note_id=?1",
                    note_id,
                ),
            },
        )
        self.assertEqual(current["source_index_status"], "ready", current)
        self.assertEqual(len(heads), 1)
        self.assertNotEqual(
            heads[0]["revision_id"],
            result["revision_id"],
            {
                "current": current,
                "revisions": self.client.portal.call(
                    self.database.fetch,
                    "SELECT id,processing_revision,source_sha256,extracted_sha256 FROM "
                    "source_revisions WHERE note_id=?1",
                    note_id,
                ),
                "jobs": self.client.portal.call(
                    self.database.fetch,
                    "SELECT processing_revision,status,last_error_code FROM "
                    "source_index_jobs WHERE note_id=?1",
                    note_id,
                ),
            },
        )

        old_chunk = self.call(
            "get",
            f"/api/classes/{self.class_id}/sources/chunks/{result['chunk_revision_id']}",
            "source_reader",
        )
        self.assertEqual(old_chunk.status_code, 200, old_chunk.text)
        self.assertEqual(old_chunk.json()["text"], original_text.split("\n", 1)[1])

        async def try_stale_write() -> str | None:
            async with self.database.acquire() as connection:
                async with connection.transaction():
                    return await persist_revision(
                        connection,
                        note_id=note_id,
                        processing_revision=note["processing_revision"],
                        source_hash=None,
                        extracted_text="stale data",
                        extraction_method="test",
                        blocks=[SourceBlock("stale", ("stale",), "stale data", {}, False, "test")],
                        mark_note_ready=False,
                    )

        stale = self.client.portal.call(try_stale_write)
        self.assertIsNone(stale)
        current_head = self.client.portal.call(
            self.database.fetchval,
            "SELECT revision_id FROM source_index_heads WHERE note_id=?1",
            note_id,
        )
        self.assertEqual(current_head, heads[0]["revision_id"])

    def test_unchanged_reprocess_rebinds_current_head_without_mutating_revision(self) -> None:
        self.add_reader("unchanged_reader")
        original_text = "# Kernel\nThe kernel coordinates CPU scheduling and memory access."
        note_id = self.seed_processing_note(original_text)
        self.run_upload_processor(note_id)
        result = self.call(
            "get",
            f"/api/classes/{self.class_id}/sources/search",
            "unchanged_reader",
            params={"query": "kernel scheduling"},
        )
        self.assertEqual(result.status_code, 200, result.text)
        chunk_revision_id = result.json()["items"][0]["chunk_revision_id"]
        before = self.client.portal.call(
            self.database.fetchrow,
            "SELECT h.revision_id,h.processing_revision,"
            "r.processing_revision AS revision_generation "
            "FROM source_index_heads h JOIN source_revisions r ON r.id=h.revision_id "
            "WHERE h.note_id=?1",
            note_id,
        )
        self.client.portal.call(
            self.database.execute,
            "UPDATE course_notes SET status='processing' WHERE id=?1",
            note_id,
        )
        processing_revision = self.client.portal.call(
            self.database.fetchval,
            "SELECT processing_revision FROM course_notes WHERE id=?1",
            note_id,
        )
        with patch("app.api.class_library.create_default_document", new_callable=AsyncMock):
            self.client.portal.call(
                _reprocess_upload_job,
                self.client.app,
                UUID(note_id),
                True,
                processing_revision,
            )
        after = self.client.portal.call(
            self.database.fetchrow,
            "SELECT h.revision_id,h.processing_revision,"
            "r.processing_revision AS revision_generation,"
            "n.processing_revision AS note_generation,n.source_index_status "
            "FROM source_index_heads h JOIN source_revisions r ON r.id=h.revision_id "
            "JOIN course_notes n ON n.id=h.note_id WHERE h.note_id=?1",
            note_id,
        )
        self.assertEqual(after["revision_id"], before["revision_id"])
        self.assertEqual(after["revision_generation"], before["revision_generation"])
        self.assertEqual(after["processing_revision"], after["note_generation"])
        self.assertEqual(after["source_index_status"], "ready")

        async def validate_bank_and_graph_refs():
            async with self.database.acquire() as connection:
                return await _validate_current_chunk_refs(
                    connection,
                    class_id=self.class_id,
                    user_id=self.owner["id"],
                    chunk_revision_ids=[chunk_revision_id],
                )

        valid_refs = self.client.portal.call(validate_bank_and_graph_refs)
        self.assertIn(chunk_revision_id, valid_refs)
        refetched = self.call(
            "get",
            f"/api/classes/{self.class_id}/sources/chunks/{chunk_revision_id}",
            "unchanged_reader",
        )
        self.assertEqual(refetched.status_code, 200, refetched.text)

    def test_draft_scheduled_and_membership_visibility_are_live(self) -> None:
        self.add_reader("draft_reader")
        note_id = self.seed_processing_note("# Draft\nDraft-only zebra fact.", state="draft")
        self.run_upload_processor(note_id)
        route = f"/api/classes/{self.class_id}/sources/search"
        owner_result = self.call("get", route, "source_owner", params={"query": "zebra"})
        self.assertEqual(owner_result.status_code, 200)
        self.assertEqual(len(owner_result.json()["items"]), 1)
        reader_result = self.call("get", route, "draft_reader", params={"query": "zebra"})
        self.assertEqual(reader_result.status_code, 200)
        self.assertEqual(reader_result.json()["items"], [])
        draft_chunk_id = owner_result.json()["items"][0]["chunk_revision_id"]
        hidden_fetch = self.call(
            "get",
            f"/api/classes/{self.class_id}/sources/chunks/{draft_chunk_id}",
            "draft_reader",
        )
        self.assertEqual(hidden_fetch.status_code, 404)

        self.client.portal.call(
            self.database.execute,
            "UPDATE course_notes SET "
            "publication_state='scheduled',release_at='2099-01-01T00:00:00.000Z' "
            "WHERE id=?1",
            note_id,
        )
        scheduled_result = self.call("get", route, "draft_reader", params={"query": "zebra"})
        self.assertEqual(scheduled_result.json()["items"], [])
        self.client.portal.call(
            self.database.execute,
            "UPDATE course_notes SET publication_state='published',release_at=NULL WHERE id=?1",
            note_id,
        )
        published_result = self.call("get", route, "draft_reader", params={"query": "zebra"})
        self.assertEqual(len(published_result.json()["items"]), 1)
        published_chunk_id = published_result.json()["items"][0]["chunk_revision_id"]

        removed = self.call(
            "delete",
            f"/api/classes/{self.class_id}/members/{self.users['draft_reader']['id']}",
            "source_owner",
        )
        self.assertEqual(removed.status_code, 204, removed.text)
        revoked_search = self.call("get", route, "draft_reader", params={"query": "zebra"})
        self.assertEqual(revoked_search.status_code, 404)
        revoked_fetch = self.call(
            "get",
            f"/api/classes/{self.class_id}/sources/chunks/{published_chunk_id}",
            "draft_reader",
        )
        self.assertEqual(revoked_fetch.status_code, 404)

    def test_assessment_answer_key_release_and_private_attempt_filter(self) -> None:
        self.add_reader("assessment_reader")
        note_id = self.seed_processing_note("# Exam\nanswerkeysecret: option c", state="published")
        self.run_upload_processor(note_id)
        assessment_id = str(uuid4())
        self.client.portal.call(
            self.database.execute,
            "INSERT INTO uploaded_assessments "
            "(id,class_id,note_id,owner_id,answer_key_release_state,"
            "processing_status,has_student_responses,questions,grading_mode) "
            "VALUES (?1,?2,?3,?4,'draft','ready',0,'[]','not_graded')",
            assessment_id,
            self.class_id,
            note_id,
            self.owner["id"],
        )
        route = f"/api/classes/{self.class_id}/sources/search"
        student_hidden = self.call(
            "get", route, "assessment_reader", params={"query": "answerkeysecret"}
        )
        self.assertEqual(student_hidden.json()["items"], [])
        student_upload = self.call(
            "get",
            route,
            "assessment_reader",
            params={"query": "answerkeysecret", "upload_id": note_id},
        )
        self.assertEqual(student_upload.status_code, 404)
        manager_visible = self.call(
            "get", route, "source_owner", params={"query": "answerkeysecret"}
        )
        self.assertEqual(len(manager_visible.json()["items"]), 1)

        self.client.portal.call(
            self.database.execute,
            "UPDATE uploaded_assessments SET answer_key_release_state='published' WHERE id=?1",
            assessment_id,
        )
        student_released = self.call(
            "get", route, "assessment_reader", params={"query": "answerkeysecret"}
        )
        self.assertEqual(len(student_released.json()["items"]), 1)

        self.client.portal.call(
            self.database.execute,
            "UPDATE uploaded_assessments SET "
            "processing_status='needs_review',has_student_responses=1,"
            "grading_mode='needs_review',answer_key_release_state='published' WHERE id=?1",
            assessment_id,
        )
        private_assessment = self.call(
            "get", route, "source_owner", params={"query": "answerkeysecret", "upload_id": note_id}
        )
        self.assertEqual(private_assessment.status_code, 404)
        private_result = self.call(
            "get", route, "source_owner", params={"query": "answerkeysecret"}
        )
        self.assertEqual(private_result.json()["items"], [])

    def test_index_only_retry_works_without_provider_and_preserves_extraction(self) -> None:
        note_id = self.seed_processing_note("# Recovery\nKeep this extracted source exactly.")
        self.run_upload_processor(note_id)
        original_text = self.client.portal.call(
            self.database.fetchval,
            "SELECT extracted_text FROM course_notes WHERE id=?1",
            note_id,
        )
        self.client.portal.call(
            self.database.execute,
            "UPDATE course_notes SET source_index_status='failed' WHERE id=?1",
            note_id,
        )
        self.client.portal.call(
            self.database.execute,
            "UPDATE source_index_jobs SET "
            "status='failed',attempts=max_attempts,last_error_code='index_failed' "
            ""
            "WHERE note_id=?1",
            note_id,
        )
        self.client.app.state.gemini_content_service = None
        response = self.call("post", f"/api/uploads/{note_id}/retry", "source_owner")
        self.assertEqual(response.status_code, 202, response.text)
        current = self.client.portal.call(
            self.database.fetchrow,
            "SELECT status,source_index_status,extracted_text FROM course_notes WHERE id=?1",
            note_id,
        )
        self.assertEqual(current["status"], "ready")
        self.assertEqual(current["source_index_status"], "ready")
        self.assertEqual(current["extracted_text"], original_text)


class SourceStructureTests(unittest.TestCase):
    def test_transcript_segments_are_coalesced_with_time_range(self) -> None:
        segments = [
            {
                "start_seconds": index * 2,
                "end_seconds": index * 2 + 1.5,
                "text": "biology facts " * 30,
            }
            for index in range(12)
        ]
        blocks = make_source_blocks(
            path=None,
            kind="lecture_audio",
            source_type="file",
            mime_type="audio/mpeg",
            extracted_text="",
            transcript_segments=segments,
        )
        self.assertLess(len(blocks), len(segments))
        self.assertEqual(blocks[0].locator["timestamp_start_seconds"], 0)
        self.assertGreater(blocks[0].locator["timestamp_end_seconds"], 1.5)
        self.assertLessEqual(len(blocks[0].text.split()), 600)

    def test_docx_interleaves_tables_and_paragraphs_and_pptx_marks_generated_descriptions(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as directory:
            docx_path = Path(directory) / "ordered.docx"
            doc = Document()
            doc.add_heading("Cell Biology", level=1)
            doc.add_paragraph("First paragraph")
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Mitochondria"
            table.cell(0, 1).text = "ATP"
            doc.add_paragraph("Last paragraph")
            doc.save(docx_path)
            doc_blocks = make_source_blocks(
                path=docx_path,
                kind="document",
                source_type="file",
                mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                extracted_text="",
            )
            text = "\n".join(block.text for block in doc_blocks)
            self.assertLess(text.index("First paragraph"), text.index("Mitochondria"))
            self.assertLess(text.index("Mitochondria"), text.index("Last paragraph"))
            self.assertIn("block_order_start", doc_blocks[0].locator)

            pptx_path = Path(directory) / "slides.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            slide.shapes.title.text = "Photosynthesis"
            textbox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
            textbox.text = "Chloroplasts convert light energy."
            presentation.save(pptx_path)
            pptx_blocks = make_source_blocks(
                path=pptx_path,
                kind="slides",
                source_type="file",
                mime_type=(
                    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
                ),
                extracted_text=(
                    "[Slide image descriptions]\n[Slide 1]\nGenerated diagram description.\n"
                    "[Slide 3]\nGenerated chart description."
                ),
            )
            generated = [
                block for block in pptx_blocks if block.provenance == "generated_description"
            ]
            self.assertEqual([block.locator for block in generated], [{"slide": 1}, {"slide": 3}])
            self.assertEqual(generated[0].extraction_method, "gemini_vision_description")
            self.assertEqual(generated[1].provenance, "generated_description")

    def test_native_pdf_splits_markdown_headings_and_keeps_page_locator(self) -> None:
        class FakePage:
            def extract_text(self) -> str:
                return (
                    "# Unit 1\nOverview of the unit.\n"
                    "## Cell membrane\nThe membrane controls material exchange.\n"
                    "### Diffusion\nParticles move down a concentration gradient."
                )

        class FakeReader:
            pages = [FakePage()]

            def __init__(self, _path: str) -> None:
                pass

        with patch("app.services.source_index.PdfReader", FakeReader):
            blocks = _pdf_blocks(Path("native.pdf"), "")

        self.assertEqual(
            [block.heading_path for block in blocks],
            [
                ("Unit 1",),
                ("Unit 1", "Cell membrane"),
                ("Unit 1", "Cell membrane", "Diffusion"),
            ],
        )
        self.assertEqual([block.locator for block in blocks], [{"page": 1}] * 3)
        self.assertEqual(blocks[1].text, "The membrane controls material exchange.")
        self.assertTrue(all(block.extraction_method == "pdf_text" for block in blocks))

    def test_pdf_page_locator_and_pasted_text_sections(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            pdf_path = Path(directory) / "scanned.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=612, height=792)
            writer.write(pdf_path)
            pdf_blocks = make_source_blocks(
                path=pdf_path,
                kind="document",
                source_type="file",
                mime_type="application/pdf",
                extracted_text="[Page 1]\nScanned page extract with enough content to identify its "
                "locator.",
            )
            self.assertEqual(pdf_blocks[0].locator["page"], 1)
            self.assertTrue(pdf_blocks[0].inferred)
        text_blocks = make_source_blocks(
            path=None,
            kind="document",
            source_type="pasted_text",
            mime_type=None,
            extracted_text="# Unit 1\nA detailed source paragraph.",
        )
        self.assertEqual(text_blocks[0].heading_path, ("Unit 1",))


class SourceStructureEdgeCaseTests(unittest.TestCase):
    def test_long_chunks_preserve_equations_code_and_exact_source_characters(self) -> None:
        source = (
            r"Equation: $\\alpha_i = x_{i+1}/(n-1)$; code: `foo.bar(x) -> y` "
            + " ".join(f"token{i}" for i in range(1100))
            + "\n\n  Preserve spacing around punctuation: a+b != a + b.  "
        )
        chunks = _split_block(SourceBlock("Notes", ("Notes",), source, {}, False, "text"))
        self.assertGreater(len(chunks), 1)
        self.assertEqual("".join(chunk for chunk, _count in chunks), source.strip())
        combined = "".join(chunk for chunk, _count in chunks)
        self.assertIn(r"$\\alpha_i = x_{i+1}/(n-1)$", combined)
        self.assertIn("`foo.bar(x) -> y`", combined)
        self.assertIn("a+b != a + b", combined)

    def test_mixed_pdf_uses_gemini_page_for_scanned_page_and_native_text_elsewhere(self) -> None:
        class FakePage:
            def __init__(self, text: str) -> None:
                self.text = text

            def extract_text(self) -> str:
                return self.text

        class FakeReader:
            pages = [FakePage(""), FakePage("Native selectable text. " * 8)]

            def __init__(self, _path: str) -> None:
                pass

        with patch("app.services.source_index.PdfReader", FakeReader):
            blocks = _pdf_blocks(
                Path("mixed.pdf"),
                "[Page 1]\nGemini recovered the scanned diagram labels and explanation.\n"
                "[Page 2]\nGemini page two fallback should not replace native text.",
            )
        self.assertEqual(len(blocks), 2)
        self.assertEqual(blocks[0].locator, {"page": 1})
        self.assertEqual(blocks[0].extraction_method, "gemini_pdf_vision")
        self.assertTrue(blocks[0].inferred)
        self.assertIn("Gemini recovered", blocks[0].text)
        self.assertEqual(blocks[1].extraction_method, "pdf_text")
        self.assertFalse(blocks[1].inferred)
        self.assertIn("Native selectable text", blocks[1].text)


class SourceIndexMigrationTests(unittest.IsolatedAsyncioTestCase):
    async def test_legacy_database_gets_index_status_column_and_fts5_without_data_loss(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = str(Path(directory) / "existing.sqlite")
            settings = Settings(
                _env_file=None,
                gemini_api_key="unused",
                file_url_signing_secret="test-signing-secret-with-at-least-32-characters",
                sqlite_path=path,
                lean_tool_enabled=False,
            )
            database = await create_database(settings)
            await database.execute(
                "INSERT INTO users (id,username,display_name,password_hash) "
                "VALUES ('legacy-owner','legacyowner','Legacy Owner','hash')"
            )
            await database.execute(
                "INSERT INTO course_classes (id,owner_id,name) VALUES "
                "('legacy-class','legacy-owner','Legacy')"
            )
            await database.execute(
                "INSERT INTO course_notes "
                "(id,class_id,title,kind,source_type,extracted_text,status,processing_progress) "
                "VALUES ('legacy-note','legacy-class','Legacy "
                "notes','document','pasted_text','kept text','ready',100)"
            )
            await database.close()
            raw = sqlite3.connect(path)
            raw.execute("DROP TRIGGER IF EXISTS concept_extraction_enqueue_head_insert")
            raw.execute("DROP TRIGGER IF EXISTS concept_extraction_enqueue_head_update")
            raw.execute("DROP TRIGGER IF EXISTS concept_extraction_enqueue_note_ready")
            raw.execute("ALTER TABLE course_notes DROP COLUMN source_index_status")
            raw.commit()
            raw.close()

            migrated = await create_database(settings)
            self.addAsyncCleanup(migrated.close)
            self.assertEqual(
                await migrated.fetchval(
                    "SELECT source_index_status FROM course_notes WHERE id='legacy-note'"
                ),
                "unindexed",
            )
            self.assertEqual(
                await migrated.fetchval(
                    "SELECT extracted_text FROM course_notes WHERE id='legacy-note'"
                ),
                "kept text",
            )
            self.assertEqual(
                await migrated.fetchval(
                    "SELECT count(*) FROM sqlite_master WHERE type='table' AND "
                    "name='source_chunks_fts'"
                ),
                1,
            )
            self.assertEqual(await migrated.fetch("PRAGMA foreign_key_check"), [])
