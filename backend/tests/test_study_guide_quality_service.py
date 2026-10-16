"""Model-assisted study-guide review, repair, and legacy conversion tests."""

from __future__ import annotations

import io
import json
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch
from uuid import uuid4

from google.genai import types
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter

from app.api.schemas import StudyGuideSectionV2
from app.config import Settings
from app.services.gemini_content import (
    ContentGenerationFailure,
    GeminiContentService,
    VisualEvidenceUnavailableError,
    _StudyGuideCitationLocatorRepair,
    _StudyGuideCitationLocatorRepairs,
    _StudyGuideEvidenceQuote,
    _StudyGuideFindingAdjudication,
    _StudyGuideFindingDisposition,
    _StudyGuideRepairBatch,
    _StudyGuideRepairEdit,
    _StudyGuideRepairEdits,
    _StudyGuideReview,
    _StudyGuideReviewFinding,
)


def settings() -> Settings:
    return Settings(
        _env_file=None,
        sqlite_path=":memory:",
        file_url_signing_secret="test-signing-secret-with-at-least-32-characters",
    )


def source(upload_id) -> dict:
    return {
        "id": str(upload_id),
        "title": "Mechanics notes",
        "kind": "notes",
        "extracted_text": "Traction is the stress vector acting on a plane.",
    }


def section(upload_id, *, heading: str = "Traction", text: str = "Traction is a vector."):
    return StudyGuideSectionV2.model_validate(
        {
            "heading": heading,
            "key_terms": ["traction"],
            "sources": [{"upload_id": str(upload_id), "page": 2}],
            "blocks": [{"type": "paragraph", "text": text}],
        }
    )


class StudyGuideQualityServiceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.upload_id = uuid4()
        self.sources = [source(self.upload_id)]
        self.service = GeminiContentService(None, settings())

    def test_review_sends_cited_source_and_returns_specific_model_findings(self) -> None:
        self.service._generate_json = Mock(
            return_value=_StudyGuideReview(
                approved=False,
                findings=[
                    _StudyGuideReviewFinding(
                        category="calculation",
                        location="sections[0].blocks[0].text",
                        finding="Recompute the projected traction value.",
                    )
                ],
            )
        )
        review_section = section(self.upload_id).model_dump(mode="json")
        review_section["blocks"].append(
            {
                "type": "labeled_diagram",
                "template_id": "molecule_2d_v1",
                "title": "Water molecule",
                "parameters": {"molecule_id": "water"},
                "sources": [{"upload_id": str(self.upload_id), "page": 2}],
            }
        )
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [review_section],
        }

        findings = self.service.review_study_guide(candidate, self.sources)

        self.assertEqual(findings[0]["category"], "calculation")
        prompt = self.service._generate_json.call_args.args[0]
        self.assertIn("Traction is the stress vector", prompt)
        self.assertIn("Recompute every shown calculation", prompt)
        self.assertIn("do not require every citation to support every claim", prompt)
        self.assertIn(
            "A worked example that explicitly applies and verifies an equation counts",
            prompt,
        )
        self.assertIn("at most 2 highest-priority findings", prompt)
        self.assertIn("at most 30 words", prompt)
        self.assertIn("a citation, ID, structural object", prompt)
        self.assertNotIn('"body":', prompt)
        self.assertNotIn('"fallback_text":', prompt)
        self.assertNotIn('"alt_text":', prompt)
        self.assertEqual(
            self.service._generate_json.call_args.kwargs["max_output_tokens"], 3_000
        )
        self.assertEqual(self.service._generate_json.call_args.kwargs["temperature"], 0.0)

    def test_review_attaches_images_from_only_the_cited_powerpoint_slide(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            image_paths = []
            presentation_path = f"{directory}/slides.pptx"
            for index, color in enumerate(("red", "blue")):
                image_path = f"{directory}/slide-{index + 1}.png"
                Image.new("RGB", (32, 32), color).save(image_path)
                image_paths.append(image_path)
            presentation = Presentation()
            for image_path in image_paths:
                slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                slide.shapes.add_picture(
                    image_path, Inches(1), Inches(1), width=Inches(2), height=Inches(2)
                )
            presentation.save(presentation_path)

            visual_source = source(self.upload_id)
            visual_source.update(
                {
                    "extracted_text": (
                        "[Slide 1]\nFirst slide text.\n\n"
                        "[Slide 2]\nSecond slide text."
                    ),
                    "original_path": presentation_path,
                }
            )
            candidate = {
                "format": "study_guide",
                "schema_version": 2,
                "sections": [
                    {
                        "heading": "Principal stress",
                        "key_terms": ["principal stress"],
                        "sources": [{"upload_id": str(self.upload_id), "slide": 1}],
                        "blocks": [
                            {
                                "type": "paragraph",
                                "text": "A principal plane carries no shear traction.",
                            }
                        ],
                    }
                ],
            }
            self.service._generate_json = Mock(
                return_value=_StudyGuideReview(approved=True, findings=[])
            )

            self.service.review_study_guide(candidate, [visual_source])

            request_content = self.service._generate_json.call_args.args[0]
            self.assertIsInstance(request_content, list)
            text_parts = [item for item in request_content if isinstance(item, str)]
            image_parts = [
                item
                for item in request_content
                if isinstance(item, types.Part) and item.inline_data is not None
            ]
            self.assertIn("Visual evidence 1: Mechanics notes, slide 1", "".join(text_parts))
            self.assertEqual(len(image_parts), 1)
            self.assertEqual(image_parts[0].inline_data.mime_type, "image/png")
            with Image.open(io.BytesIO(image_parts[0].inline_data.data)) as image:
                if shutil.which("soffice") and shutil.which("pdftoppm"):
                    self.assertGreater(image.width, 1000)
                    self.assertEqual(image.getpixel((200, 200)), (255, 0, 0))
                else:
                    self.assertEqual(image.getpixel((0, 0)), (255, 0, 0))

            self.service._generate_json = Mock(
                return_value=_StudyGuideRepairEdits(
                    edits=[
                        _StudyGuideRepairEdit(
                            location="sections[0].blocks[0].text",
                            operation="replace",
                            match="no shear traction",
                            replacement="zero shear traction",
                        )
                    ]
                )
            )
            repaired = self.service.repair_study_guide(
                candidate,
                [visual_source],
                [
                    {
                        "category": "accuracy",
                        "location": "sections[0].blocks[0].text",
                        "finding": "Correct the claim against the cited slide.",
                    }
                ],
            )
            repair_content = self.service._generate_json.call_args.args[0]
            self.assertIsInstance(repair_content, list)
            self.assertTrue(
                any(
                    isinstance(item, types.Part) and item.inline_data is not None
                    for item in repair_content
                )
            )
            self.assertEqual(
                repaired["sections"][0]["blocks"][0]["text"],
                "A principal plane carries zero shear traction.",
            )

    def test_review_attaches_embedded_image_from_the_cited_pdf_page(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            pdf_path = f"{directory}/scanned-notes.pdf"
            Image.new("RGB", (64, 64), "green").save(pdf_path, "PDF")
            visual_source = source(self.upload_id)
            visual_source.update(
                {
                    "extracted_text": "[Page 1]\nStress diagram.",
                    "original_path": pdf_path,
                }
            )
            candidate = {
                "format": "study_guide",
                "schema_version": 2,
                "sections": [
                    {
                        "heading": "Stress diagram",
                        "key_terms": ["stress"],
                        "sources": [{"upload_id": str(self.upload_id), "page": 1}],
                        "blocks": [
                            {"type": "paragraph", "text": "Stress acts on a plane."}
                        ],
                    }
                ],
            }
            self.service._generate_json = Mock(
                return_value=_StudyGuideReview(approved=True, findings=[])
            )

            self.service.review_study_guide(candidate, [visual_source])

            request_content = self.service._generate_json.call_args.args[0]
            image_parts = [
                item
                for item in request_content
                if isinstance(item, types.Part) and item.inline_data is not None
            ]
            self.assertEqual(len(image_parts), 1)
            self.assertIn(image_parts[0].inline_data.mime_type, {"image/jpeg", "image/png"})

    def test_review_can_check_selected_section_and_maps_location_to_full_guide(self) -> None:
        sections = [
            section(self.upload_id, heading=f"Section {index}").model_dump(mode="json")
            for index in range(3)
        ]
        self.service._generate_json = Mock(
            return_value=_StudyGuideReview(
                approved=False,
                findings=[
                    _StudyGuideReviewFinding(
                        category="accuracy",
                        location="sections[0].blocks[0].text",
                        finding="Clarify the selected claim.",
                    )
                ],
            )
        )

        findings = self.service.review_study_guide(
            {"format": "study_guide", "schema_version": 2, "sections": sections},
            self.sources,
            _selected_sections={2},
        )

        self.assertEqual(findings[0]["location"], "sections[2].blocks[0].text")
        prompt = self.service._generate_json.call_args.args[0]
        self.assertIn("Selected section indices in the complete guide: [2]", prompt)

    def test_adjudicator_drops_only_claims_directly_supported_by_sources(self) -> None:
        self.service._generate_json = Mock(
            side_effect=[
                _StudyGuideFindingAdjudication(
                    decisions=[
                        _StudyGuideFindingDisposition(index=0, supported=False),
                        _StudyGuideFindingDisposition(index=1, supported=False),
                    ]
                ),
                _StudyGuideFindingAdjudication(
                    decisions=[
                        _StudyGuideFindingDisposition(
                            index=0,
                            supported=True,
                            evidence_quotes=[
                                _StudyGuideEvidenceQuote(
                                    source_excerpt_index=0,
                                    quote="Traction is the stress vector acting on a plane.",
                                )
                            ],
                        )
                    ]
                ),
                _StudyGuideFindingAdjudication(
                    decisions=[
                        _StudyGuideFindingDisposition(
                            index=0,
                            supported=True,
                            evidence_quotes=[
                                _StudyGuideEvidenceQuote(
                                    source_excerpt_index=0,
                                    quote="This quote does not occur in the source.",
                                )
                            ],
                        )
                    ]
                ),
            ]
        )
        review_section = section(self.upload_id).model_dump(mode="json")
        review_section["blocks"].append(
            {
                "type": "paragraph",
                "text": "Young's modulus equals stress.",
                "sources": [{"upload_id": str(self.upload_id), "page": 2}],
            }
        )
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [review_section],
        }
        findings = [
            {
                "category": "unsupported_claim",
                "location": "sections[0].blocks[0].text",
                "finding": "The cited worked example supports this traction statement.",
            },
            {
                "category": "unsupported_claim",
                "location": "sections[0].blocks[1].text",
                "finding": "The cited source does not support this modulus statement.",
            },
        ]

        remaining, supported = self.service.adjudicate_unsupported_study_guide_findings(
            candidate, self.sources, findings
        )

        self.assertEqual(remaining, [findings[1]])
        fingerprint = self.service._study_guide_target_fingerprint(
            self.service.normalize_study_guide(candidate, self.sources),
            "sections[0].blocks[0].text",
        )
        self.assertIn(("sections[0].blocks[0].text", fingerprint), supported)
        prompt = self.service._generate_json.call_args.args[0]
        self.assertIn("Traction is the stress vector", prompt)
        self.assertIn("A worked example counts as support", prompt)
        self.assertIn("short, exact, continuous evidence quote", prompt)
        self.assertEqual(self.service._generate_json.call_args.kwargs["temperature"], 0.0)
        self.assertEqual(self.service._generate_json.call_args.kwargs["max_output_tokens"], 3_000)
        self.assertEqual(self.service._generate_json.call_count, 3)

    def test_adjudicator_fails_closed_on_invalid_supported_finding_index(self) -> None:
        self.service._generate_json = Mock(
            return_value=_StudyGuideFindingAdjudication(
                decisions=[
                    _StudyGuideFindingDisposition(
                        index=1,
                        supported=True,
                        evidence_quotes=[
                            _StudyGuideEvidenceQuote(
                                source_excerpt_index=0,
                                quote="Traction is the stress vector acting on a plane.",
                            )
                        ],
                    )
                ]
            )
        )
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [section(self.upload_id).model_dump(mode="json")],
        }
        with self.assertRaises(ContentGenerationFailure) as raised:
            self.service.adjudicate_unsupported_study_guide_findings(
                candidate,
                self.sources,
                [
                    {
                        "category": "unsupported_claim",
                        "location": "sections[0].blocks[0].text",
                        "finding": "The source does not support this statement.",
                    }
                ],
            )

        self.assertEqual(
            raised.exception.diagnostic_code, "quality_adjudication_index_invalid"
        )

    def test_adjudicator_retries_token_limited_batches_per_finding(self) -> None:
        self.service._generate_json = Mock(
            side_effect=[
                ContentGenerationFailure(
                    "The provider reached its output token limit.",
                    code="output_token_limit",
                    diagnostic_code="completion_token_limit",
                ),
                _StudyGuideFindingAdjudication(
                    decisions=[
                        _StudyGuideFindingDisposition(
                            index=0,
                            supported=True,
                            evidence_quotes=[
                                _StudyGuideEvidenceQuote(
                                    source_excerpt_index=0,
                                    quote="Traction is the stress vector acting on a plane.",
                                )
                            ],
                        )
                    ]
                ),
                _StudyGuideFindingAdjudication(
                    decisions=[_StudyGuideFindingDisposition(index=0, supported=False)]
                ),
            ]
        )
        review_section = section(self.upload_id).model_dump(mode="json")
        review_section["blocks"].append(
            {
                "type": "paragraph",
                "text": "Young's modulus equals stress.",
                "sources": [{"upload_id": str(self.upload_id), "page": 2}],
            }
        )
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [review_section],
        }
        findings = [
            {
                "category": "unsupported_claim",
                "location": f"sections[0].blocks[{index}].text",
                "finding": f"Finding {index}.",
            }
            for index in range(2)
        ]

        remaining, supported = self.service.adjudicate_unsupported_study_guide_findings(
            candidate, self.sources, findings
        )

        self.assertEqual(remaining, [findings[1]])
        self.assertEqual(len(supported), 1)
        self.assertEqual(self.service._generate_json.call_count, 3)

    def test_review_splits_token_limited_section_into_blocks_and_maps_locations(self) -> None:
        self.service._generate_json = Mock(
            side_effect=[
                ContentGenerationFailure(
                    "The provider reached its output token limit.",
                    code="output_token_limit",
                    diagnostic_code="completion_token_limit",
                ),
                _StudyGuideReview(approved=True, findings=[]),
                _StudyGuideReview(
                    approved=False,
                    findings=[
                        _StudyGuideReviewFinding(
                            category="accuracy",
                            location="sections[0].blocks[0].text",
                            finding="Clarify the force direction.",
                        )
                    ],
                ),
            ]
        )
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [
                StudyGuideSectionV2.model_validate(
                    {
                        "heading": "Traction",
                        "key_terms": ["traction"],
                        "sources": [{"upload_id": str(self.upload_id), "page": 2}],
                        "blocks": [
                            {"type": "paragraph", "text": "Traction is a vector."},
                            {
                                "type": "takeaway",
                                "text": "Normal traction is projected onto the plane normal.",
                            },
                        ],
                    }
                ).model_dump(mode="json")
            ],
        }

        findings = self.service.review_study_guide(candidate, self.sources)

        self.assertEqual(self.service._generate_json.call_count, 3)
        self.assertEqual(findings[0]["location"], "sections[0].blocks[1].text")
        first_block_prompt = self.service._generate_json.call_args_list[1].args[0]
        self.assertIn('"id": "section-1-block-1"', first_block_prompt)
        self.assertNotIn('"id": "section-1-block-2"', first_block_prompt)

    def test_review_checks_sections_individually_and_maps_locations_back(self) -> None:
        self.service._generate_json = Mock(
            side_effect=[
                _StudyGuideReview(approved=True, findings=[]),
                _StudyGuideReview(approved=True, findings=[]),
                _StudyGuideReview(
                    approved=False,
                    findings=[
                        _StudyGuideReviewFinding(
                            category="arithmetic_error",
                            location="sections[0].blocks[0].text",
                            finding="Recompute this result.",
                        )
                    ],
                ),
                _StudyGuideReview(approved=True, findings=[]),
            ]
        )
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [
                section(self.upload_id, heading=f"Topic {index}").model_dump(mode="json")
                for index in range(4)
            ],
        }

        findings = self.service.review_study_guide(candidate, self.sources)

        self.assertEqual(len(findings), 1)
        self.assertEqual(findings[0]["location"], "sections[2].blocks[0].text")
        self.assertEqual(findings[0]["category"], "other")
        self.assertEqual(self.service._generate_json.call_count, 4)
        prompts = [call.args[0] for call in self.service._generate_json.call_args_list]
        self.assertIn("Selected section indices in the complete guide: [0]", prompts[0])
        self.assertIn("Selected section indices in the complete guide: [1]", prompts[1])
        self.assertIn("Selected section indices in the complete guide: [2]", prompts[2])
        self.assertIn("Selected section indices in the complete guide: [3]", prompts[3])
        self.assertIn("compact guide context for cross-section consistency", prompts[2])

    def test_review_maps_a_nonexistent_step_leaf_to_its_authored_block(self) -> None:
        self.service._generate_json = Mock(
            return_value=_StudyGuideReview(
                approved=False,
                findings=[
                    _StudyGuideReviewFinding(
                        category="accuracy",
                        location="sections[0].blocks[1].steps[99].explanation",
                        finding="Check this step against the cited source.",
                    )
                ],
            )
        )
        section_data = section(self.upload_id).model_dump(mode="json")
        section_data["blocks"].append(
            {
                "type": "example_steps",
                "title": "Short derivation",
                "steps": [
                    {"title": "First", "explanation": "Start from equilibrium."},
                    {"title": "Second", "explanation": "Rearrange the relation."},
                ],
                "sources": [{"upload_id": str(self.upload_id), "page": 2}],
            }
        )
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [section_data],
        }

        findings = self.service.review_study_guide(candidate, self.sources)

        self.assertEqual(len(findings), 1)
        self.assertEqual(findings[0]["location"], "sections[0].blocks[1]")

    def test_repair_targets_respect_optional_math_fields_for_each_block_type(self) -> None:
        section_data = section(self.upload_id).model_dump(mode="json")
        section_data["blocks"].extend(
            [
                {
                    "type": "formula",
                    "expression": "x=1",
                    "definitions": [{"symbol": "x", "meaning": "value", "unit": None}],
                    "sources": [{"upload_id": str(self.upload_id), "page": 2}],
                },
                {
                    "type": "example_steps",
                    "title": "Derivation",
                    "expression": None,
                    "steps": [
                        {"title": "First", "explanation": "Start here.", "expression": None},
                        {"title": "Second", "explanation": "Continue here."},
                    ],
                    "sources": [{"upload_id": str(self.upload_id), "page": 2}],
                },
            ]
        )
        content = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [section_data],
        }

        editable = self.service._study_guide_editable_values(content)

        self.assertNotIn("sections[0].blocks[0].expression", editable)
        self.assertIn("sections[0].blocks[1].expression", editable)
        self.assertIn("sections[0].blocks[1].definitions[0].unit", editable)
        self.assertNotIn("sections[0].blocks[2].expression", editable)
        self.assertIn("sections[0].blocks[2].steps[0].expression", editable)

    def test_review_falls_back_to_blocks_for_a_truncated_section(self) -> None:
        self.service._generate_json = Mock(
            side_effect=[
                ContentGenerationFailure(
                    "The provider reached its output token limit.",
                    code="output_token_limit",
                    diagnostic_code="completion_token_limit",
                ),
                _StudyGuideReview(approved=True, findings=[]),
                _StudyGuideReview(
                    approved=False,
                    findings=[
                        _StudyGuideReviewFinding(
                            category="accuracy",
                            location="sections[0].blocks[0].text",
                            finding="Clarify the force direction.",
                        )
                    ],
                ),
                _StudyGuideReview(approved=True, findings=[]),
            ]
        )
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [
                section(self.upload_id, heading=f"Topic {index}").model_dump(mode="json")
                for index in range(2)
            ],
        }

        findings = self.service.review_study_guide(candidate, self.sources)

        self.assertEqual(self.service._generate_json.call_count, 3)
        self.assertEqual(findings[0]["location"], "sections[1].blocks[0].text")
        prompts = [call.args[0] for call in self.service._generate_json.call_args_list]
        self.assertIn("Selected section indices in the complete guide: [0]", prompts[0])
        self.assertIn('"heading": "Topic 0"', prompts[1])
        self.assertIn("Selected section indices in the complete guide: [1]", prompts[2])

    def test_truncated_block_gets_one_compact_review_retry(self) -> None:
        token_limit = ContentGenerationFailure(
            "The provider reached its output token limit.",
            code="output_token_limit",
            diagnostic_code="completion_token_limit",
        )
        self.service._generate_json = Mock(
            side_effect=[
                token_limit,
                token_limit,
                _StudyGuideReview(
                    approved=False,
                    findings=[
                        _StudyGuideReviewFinding(
                            category="accuracy",
                            location="sections[0].blocks[0].text",
                            finding="Clarify the force direction.",
                        )
                    ],
                ),
            ]
        )
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [section(self.upload_id).model_dump(mode="json")],
        }

        findings = self.service.review_study_guide(candidate, self.sources)

        self.assertEqual(self.service._generate_json.call_count, 3)
        self.assertEqual(findings[0]["location"], "sections[0].blocks[0].text")
        self.assertEqual(
            [
                call.kwargs["max_output_tokens"]
                for call in self.service._generate_json.call_args_list
            ],
            [3_000, 3_000, 1_500],
        )
        compact_prompt = self.service._generate_json.call_args_list[2].args[0]
        self.assertIn("at most 2 highest-priority findings", compact_prompt)
        self.assertIn("at most 25 words", compact_prompt)

    def test_repair_applies_only_the_exact_flagged_text_span(self) -> None:
        unchanged_block = {
            "type": "takeaway",
            "text": "Keep the separate source-supported takeaway untouched.",
            "sources": [{"upload_id": str(self.upload_id)}],
        }
        flagged_section = section(
            self.upload_id,
            heading="Stress tensor",
            text="An inaccurate old claim. Keep this supported explanation.",
        ).model_dump(mode="json")
        flagged_section["blocks"].append(unchanged_block)
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [
                section(self.upload_id, text="Traction is a vector.").model_dump(mode="json"),
                flagged_section,
            ],
        }
        self.service._generate_json = Mock(
            return_value=_StudyGuideRepairEdits(
                edits=[
                    _StudyGuideRepairEdit(
                        location="sections[1].blocks[0].text",
                        operation="replace",
                        match="An inaccurate old claim.",
                        replacement="A tensor encodes the stress state at a point.",
                    )
                ]
            )
        )
        original = self.service._normalized_quality_guide(candidate, self.sources)

        result = self.service.repair_study_guide(
            candidate,
            self.sources,
            [
                {
                    "category": "accuracy",
                    "location": "sections[1].blocks[0].text",
                    "finding": "Correct this claim against the cited material.",
                }
            ],
        )

        self.assertEqual(result["sections"][0], original["sections"][0])
        self.assertEqual(result["sections"][1]["heading"], "Stress tensor")
        self.assertEqual(
            result["sections"][1]["blocks"][0]["text"],
            "A tensor encodes the stress state at a point. Keep this supported explanation.",
        )
        self.assertEqual(
            result["sections"][1]["blocks"][0]["sources"],
            original["sections"][1]["blocks"][0]["sources"],
        )
        self.assertEqual(
            result["sections"][1]["blocks"][1],
            original["sections"][1]["blocks"][1],
        )
        self.assertEqual(result["sections"][1]["key_terms"], original["sections"][1]["key_terms"])
        prompt = self.service._generate_json.call_args.args[0]
        self.assertIn('"location": "sections[1].blocks[0].text"', prompt)
        self.assertIn("An inaccurate old claim", prompt)
        self.assertNotIn('"heading": "Stress tensor"', prompt)
        self.assertIn("an edits array", prompt)
        self.assertNotIn('"sections":', prompt)

    def test_repair_retries_a_noop_with_a_specific_in_place_correction(self) -> None:
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [
                section(
                    self.upload_id,
                    text="An unsupported claim. Keep this supported explanation.",
                ).model_dump(mode="json")
            ],
        }
        self.service._generate_json = Mock(
            side_effect=[
                _StudyGuideRepairEdits(edits=[]),
                _StudyGuideRepairEdits(
                    edits=[
                        _StudyGuideRepairEdit(
                            location="sections[0].blocks[0].text",
                            operation="replace",
                            match="An unsupported claim.",
                            replacement="The cited course excerpt does not establish this detail.",
                        )
                    ]
                ),
            ]
        )

        result = self.service.repair_study_guide(
            candidate,
            self.sources,
            [
                {
                    "category": "unsupported_claim",
                    "location": "sections[0].blocks[0].text",
                    "finding": "This claim is not established by the cited excerpt.",
                }
            ],
        )

        self.assertEqual(self.service._generate_json.call_count, 2)
        retry_prompt = self.service._generate_json.call_args_list[1].args[0]
        self.assertIn("previous repair did not change any flagged field", retry_prompt)
        self.assertEqual(
            result["sections"][0]["blocks"][0]["text"],
            "The cited course excerpt does not establish this detail. "
            "Keep this supported explanation.",
        )

    def test_grouped_repair_retries_when_no_field_changes(self) -> None:
        candidate_section = section(
            self.upload_id, text="The first unsupported statement."
        ).model_dump(mode="json")
        candidate_section["blocks"].append(
            {
                "type": "takeaway",
                "text": "The second unsupported statement.",
                "sources": [{"upload_id": str(self.upload_id), "page": 2}],
            }
        )
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [candidate_section],
        }
        first = "sections[0].blocks[0].text"
        second = "sections[0].blocks[1].text"
        self.service._generate_json = Mock(
            side_effect=[
                _StudyGuideRepairEdits(edits=[]),
                _StudyGuideRepairEdits(
                    edits=[
                        _StudyGuideRepairEdit(
                            location=first,
                            operation="replace",
                            match="The first unsupported statement.",
                            replacement="Traction is the stress vector acting on a plane.",
                        ),
                        _StudyGuideRepairEdit(
                            location=second,
                            operation="replace",
                            match="The second unsupported statement.",
                            replacement="Traction is the stress vector acting on a plane.",
                        ),
                    ]
                ),
            ]
        )

        result = self.service.repair_study_guide(
            candidate,
            self.sources,
            [
                {"category": "unsupported_claim", "location": first, "finding": "Correct."},
                {"category": "unsupported_claim", "location": second, "finding": "Correct."},
            ],
        )

        self.assertEqual(self.service._generate_json.call_count, 2)
        self.assertIn(
            "previous repair did not change any flagged field",
            self.service._generate_json.call_args_list[1].args[0],
        )
        self.assertEqual(
            result["sections"][0]["blocks"][0]["text"],
            "Traction is the stress vector acting on a plane.",
        )
        self.assertEqual(
            result["sections"][0]["blocks"][1]["text"],
            "Traction is the stress vector acting on a plane.",
        )

    def test_repair_groups_up_to_four_exact_fields_without_rebuilding_blocks(self) -> None:
        candidate_section = section(
            self.upload_id, text="Correct the first exact claim."
        ).model_dump(mode="json")
        candidate_section["blocks"].append(
            {
                "type": "takeaway",
                "text": "Correct the second exact claim.",
                "sources": [{"upload_id": str(self.upload_id), "page": 2}],
            }
        )
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [candidate_section],
        }
        first = "sections[0].blocks[0].text"
        second = "sections[0].blocks[1].text"
        self.service._generate_json = Mock(
            return_value=_StudyGuideRepairEdits(
                edits=[
                    _StudyGuideRepairEdit(
                        location=first,
                        operation="replace",
                        match="Correct the first exact claim.",
                        replacement="Keep the first source-grounded claim.",
                    ),
                    _StudyGuideRepairEdit(
                        location=second,
                        operation="replace",
                        match="Correct the second exact claim.",
                        replacement="Keep the second source-grounded claim.",
                    ),
                ]
            )
        )

        result = self.service.repair_study_guide(
            candidate,
            self.sources,
            [
                {"category": "accuracy", "location": first, "finding": "Fix the first."},
                {"category": "accuracy", "location": second, "finding": "Fix the second."},
            ],
        )

        self.assertEqual(self.service._generate_json.call_count, 1)
        target_json = self.service._generate_json.call_args.args[0].split(
            "Editable values and exact paths as JSON:\n", 1
        )[1].split("\n\nFindings for these locations as JSON:", 1)[0]
        self.assertEqual(
            {item["location"] for item in json.loads(target_json)}, {first, second}
        )
        self.assertEqual(
            result["sections"][0]["blocks"][0]["text"],
            "Keep the first source-grounded claim.",
        )
        self.assertEqual(
            result["sections"][0]["blocks"][1]["text"],
            "Keep the second source-grounded claim.",
        )

    def test_grouped_repair_discards_only_a_field_with_an_invalid_exact_span(self) -> None:
        candidate_section = section(
            self.upload_id, text="The first source-backed statement."
        ).model_dump(mode="json")
        candidate_section["blocks"].append(
            {
                "type": "takeaway",
                "text": "The second source-backed statement.",
                "sources": [{"upload_id": str(self.upload_id), "page": 2}],
            }
        )
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [candidate_section],
        }
        first = "sections[0].blocks[0].text"
        second = "sections[0].blocks[1].text"
        self.service._generate_json = Mock(
            return_value=_StudyGuideRepairEdits(
                edits=[
                    _StudyGuideRepairEdit(
                        location=first,
                        operation="replace",
                        match="The first source-backed statement.",
                        replacement="The corrected first statement.",
                    ),
                    _StudyGuideRepairEdit(
                        location=second,
                        operation="replace",
                        match="This exact span does not exist.",
                        replacement="The corrected second statement.",
                    ),
                ]
            )
        )

        result = self.service.repair_study_guide(
            candidate,
            self.sources,
            [
                {"category": "accuracy", "location": first, "finding": "Fix the first."},
                {"category": "accuracy", "location": second, "finding": "Fix the second."},
            ],
        )

        self.assertEqual(
            result["sections"][0]["blocks"][0]["text"],
            "The corrected first statement.",
        )
        self.assertEqual(
            result["sections"][0]["blocks"][1]["text"],
            "The second source-backed statement.",
        )

    def test_grouped_repair_keeps_valid_field_when_neighbor_edit_is_a_noop(self) -> None:
        candidate_section = section(
            self.upload_id, text="The first source-backed statement."
        ).model_dump(mode="json")
        candidate_section["blocks"].append(
            {
                "type": "takeaway",
                "text": "The second source-backed statement.",
                "sources": [{"upload_id": str(self.upload_id), "page": 2}],
            }
        )
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [candidate_section],
        }
        first = "sections[0].blocks[0].text"
        second = "sections[0].blocks[1].text"
        self.service._generate_json = Mock(
            return_value=_StudyGuideRepairEdits(
                edits=[
                    _StudyGuideRepairEdit(
                        location=first,
                        operation="replace",
                        match="The first source-backed statement.",
                        replacement="The corrected first statement.",
                    ),
                    _StudyGuideRepairEdit(
                        location=second,
                        operation="replace",
                        match="The second source-backed statement.",
                        replacement="The second source-backed statement.",
                    ),
                ]
            )
        )

        result = self.service.repair_study_guide(
            candidate,
            self.sources,
            [
                {"category": "accuracy", "location": first, "finding": "Fix the first."},
                {"category": "accuracy", "location": second, "finding": "Fix the second."},
            ],
        )

        self.assertEqual(
            result["sections"][0]["blocks"][0]["text"],
            "The corrected first statement.",
        )
        self.assertEqual(
            result["sections"][0]["blocks"][1]["text"],
            "The second source-backed statement.",
        )

    def test_single_field_noop_is_retried_with_an_exact_span_correction(self) -> None:
        text_value = "The source-supported statement."
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [section(self.upload_id, text=text_value).model_dump(mode="json")],
        }
        self.service._generate_json = Mock(
            side_effect=[
                _StudyGuideRepairEdits(
                    edits=[
                        _StudyGuideRepairEdit(
                            location="sections[0].blocks[0].text",
                            operation="replace",
                            match=text_value,
                            replacement=text_value,
                        )
                    ]
                ),
                _StudyGuideRepairEdits(
                    edits=[
                        _StudyGuideRepairEdit(
                            location="sections[0].blocks[0].text",
                            operation="replace",
                            match=text_value,
                            replacement="Traction is the stress vector acting on a plane.",
                        )
                    ]
                ),
            ]
        )

        result = self.service.repair_study_guide(
            candidate,
            self.sources,
            [
                {
                    "category": "accuracy",
                    "location": "sections[0].blocks[0].text",
                    "finding": "Review this statement again against the source.",
                }
            ],
        )

        self.assertEqual(
            result["sections"][0]["blocks"][0]["text"],
            "Traction is the stress vector acting on a plane.",
        )
        self.assertEqual(self.service._generate_json.call_count, 2)

    def test_repair_retries_one_invalid_patch_at_the_same_field(self) -> None:
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [section(self.upload_id).model_dump(mode="json")],
        }
        self.service._generate_json = Mock(
            side_effect=[
                _StudyGuideRepairEdits(
                    edits=[
                        _StudyGuideRepairEdit(
                            location="sections[0].blocks[0].text",
                            operation="replace",
                            match="This phrase is not present",
                            replacement="Traction acts on a plane",
                        )
                    ]
                ),
                _StudyGuideRepairEdits(
                    edits=[
                        _StudyGuideRepairEdit(
                            location="sections[0].blocks[0].text",
                            operation="replace",
                            match="Traction is a vector.",
                            replacement="Traction is the stress vector acting on a plane.",
                        )
                    ]
                ),
            ]
        )

        result = self.service.repair_study_guide(
            candidate,
            self.sources,
            [
                {
                    "category": "accuracy",
                    "location": "sections[0].blocks[0].text",
                    "finding": "Correct this claim against the cited material.",
                }
            ],
        )

        self.assertEqual(
            result["sections"][0]["blocks"][0]["text"],
            "Traction is the stress vector acting on a plane.",
        )
        self.assertEqual(self.service._generate_json.call_count, 2)
        self.assertIn("failed exact-span validation", self.service._generate_json.call_args.args[0])

    def test_repair_can_replace_one_key_term_value_in_place(self) -> None:
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [section(self.upload_id).model_dump(mode="json")],
        }
        location = "sections[0].key_terms[0]"
        editable = self.service._study_guide_editable_values(candidate)

        self.service._apply_study_guide_edits(
            candidate,
            editable,
            [location],
            [
                _StudyGuideRepairEdit(
                    location=location,
                    operation="replace_value",
                    match="",
                    replacement="surface force per unit area",
                )
            ],
        )

        self.assertEqual(
            candidate["sections"][0]["key_terms"], ["surface force per unit area"]
        )

    def test_key_term_replacement_recovers_when_old_value_does_not_match_exactly(self) -> None:
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [section(self.upload_id).model_dump(mode="json")],
        }
        location = "sections[0].key_terms[0]"
        editable = self.service._study_guide_editable_values(candidate)

        self.service._apply_study_guide_edits(
            candidate,
            editable,
            [location],
            [
                _StudyGuideRepairEdit(
                    location=location,
                    operation="replace",
                    match="traction with a nonmatching normalization",
                    replacement="surface force per unit area",
                )
            ],
        )

        self.assertEqual(
            candidate["sections"][0]["key_terms"], ["surface force per unit area"]
        )

    def test_key_term_noop_retry_uses_a_targeted_whole_value_edit(self) -> None:
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [section(self.upload_id).model_dump(mode="json")],
        }
        location = "sections[0].key_terms[0]"
        self.service._generate_json = Mock(
            side_effect=[
                _StudyGuideRepairEdits(
                    edits=[
                        _StudyGuideRepairEdit(
                            location=location,
                            operation="replace",
                            match="traction",
                            replacement="traction",
                        )
                    ]
                ),
                _StudyGuideRepairEdits(
                    edits=[
                        _StudyGuideRepairEdit(
                            location=location,
                            operation="replace_value",
                            match="",
                            replacement="surface traction",
                        )
                    ]
                ),
            ]
        )

        result = self.service.repair_study_guide(
            candidate,
            self.sources,
            [
                {
                    "category": "unsupported_claim",
                    "location": location,
                    "finding": "Use a source-supported term.",
                }
            ],
        )

        self.assertEqual(result["sections"][0]["key_terms"], ["surface traction"])
        self.assertEqual(self.service._generate_json.call_count, 2)
        self.assertIn("left this key term unchanged", self.service._generate_json.call_args.args[0])

    def test_key_term_whole_value_edit_is_rejected_for_other_fields(self) -> None:
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [section(self.upload_id).model_dump(mode="json")],
        }
        location = "sections[0].blocks[0].text"
        editable = self.service._study_guide_editable_values(candidate)

        with self.assertRaises(ContentGenerationFailure) as raised:
            self.service._apply_study_guide_edits(
                candidate,
                editable,
                [location],
                [
                    _StudyGuideRepairEdit(
                        location=location,
                        operation="replace_value",
                        match="",
                        replacement="A different paragraph.",
                    )
                ],
            )

        self.assertEqual(raised.exception.diagnostic_code, "quality_repair_invalid_patch")

    def test_grouped_repair_retries_when_no_exact_span_edit_can_be_applied(self) -> None:
        candidate_section = section(
            self.upload_id, text="The first claim."
        ).model_dump(mode="json")
        candidate_section["blocks"].append(
            {
                "type": "takeaway",
                "text": "The second claim.",
                "sources": [{"upload_id": str(self.upload_id), "page": 2}],
            }
        )
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [candidate_section],
        }
        first = "sections[0].blocks[0].text"
        second = "sections[0].blocks[1].text"
        self.service._generate_json = Mock(
            side_effect=[
                _StudyGuideRepairEdits(
                    edits=[
                        _StudyGuideRepairEdit(
                            location=first,
                            operation="replace",
                            match="This phrase is not present",
                            replacement="Traction acts on a plane.",
                        )
                    ]
                ),
                _StudyGuideRepairEdits(
                    edits=[
                        _StudyGuideRepairEdit(
                            location=first,
                            operation="replace",
                            match="The first claim.",
                            replacement="Traction is the stress vector acting on a plane.",
                        ),
                        _StudyGuideRepairEdit(
                            location=second,
                            operation="replace",
                            match="The second claim.",
                            replacement="Traction is the stress vector acting on a plane.",
                        ),
                    ]
                ),
            ]
        )

        result = self.service.repair_study_guide(
            candidate,
            self.sources,
            [
                {"category": "accuracy", "location": first, "finding": "Correct first."},
                {"category": "accuracy", "location": second, "finding": "Correct second."},
            ],
        )

        self.assertEqual(self.service._generate_json.call_count, 2)
        retry_prompt = self.service._generate_json.call_args_list[1].args[0]
        self.assertIn("each supplied location", retry_prompt)
        self.assertEqual(
            result["sections"][0]["blocks"][0]["text"],
            "Traction is the stress vector acting on a plane.",
        )
        self.assertEqual(
            result["sections"][0]["blocks"][1]["text"],
            "Traction is the stress vector acting on a plane.",
        )

    def test_repeated_invalid_patch_reports_only_safe_location_diagnostics(self) -> None:
        location = "sections[0].blocks[0].text"
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [section(self.upload_id).model_dump(mode="json")],
        }
        invalid_response = _StudyGuideRepairEdits(
            edits=[
                _StudyGuideRepairEdit(
                    location=location,
                    operation="replace",
                    match="Private unmatched model text.",
                    replacement="Private replacement text.",
                )
            ]
        )
        self.service._generate_json = Mock(return_value=invalid_response)

        with self.assertRaises(ContentGenerationFailure) as raised:
            self.service.repair_study_guide(
                candidate,
                self.sources,
                [
                    {
                        "category": "accuracy",
                        "location": location,
                        "finding": "Do not log this generated finding text.",
                    }
                ],
            )

        self.assertEqual(raised.exception.diagnostic_code, "quality_repair_invalid_patch")
        self.assertEqual(
            raised.exception.validation_issues,
            [{"location": location, "type": "quality_repair_invalid_patch"}],
        )
        self.assertNotIn("Private", repr(raised.exception.validation_issues))

    def test_repair_retries_one_malformed_edit_response(self) -> None:
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [section(self.upload_id).model_dump(mode="json")],
        }
        self.service._generate_json = Mock(
            side_effect=[
                ContentGenerationFailure(
                    "The configured provider returned a result that failed validation.",
                    code="invalid_model_response",
                    diagnostic_code="schema_validation_failed",
                ),
                _StudyGuideRepairEdits(
                    edits=[
                        _StudyGuideRepairEdit(
                            location="sections[0].blocks[0].text",
                            operation="replace",
                            match="Traction is a vector.",
                            replacement="Traction is the stress vector acting on a plane.",
                        )
                    ]
                ),
            ]
        )

        result = self.service.repair_study_guide(
            candidate,
            self.sources,
            [
                {
                    "category": "accuracy",
                    "location": "sections[0].blocks[0].text",
                    "finding": "Correct this claim against the cited material.",
                }
            ],
        )

        self.assertEqual(self.service._generate_json.call_count, 2)
        self.assertIn(
            "valid JSON object with an edits array",
            self.service._generate_json.call_args.args[0],
        )
        self.assertEqual(
            result["sections"][0]["blocks"][0]["text"],
            "Traction is the stress vector acting on a plane.",
        )

    def test_repair_can_fill_one_null_optional_expression_in_place(self) -> None:
        candidate_section = section(self.upload_id).model_dump(mode="json")
        candidate_section["blocks"].append(
            {
                "type": "example_steps",
                "title": "Short derivation",
                "steps": [
                    {
                        "title": "First",
                        "explanation": "Start from equilibrium.",
                        "expression": None,
                    },
                    {"title": "Second", "explanation": "Rearrange the relation."},
                ],
                "sources": [{"upload_id": str(self.upload_id), "page": 2}],
            }
        )
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [candidate_section],
        }
        self.service._generate_json = Mock(
            return_value=_StudyGuideRepairEdits(
                edits=[
                    _StudyGuideRepairEdit(
                        location="sections[0].blocks[1].steps[0].expression",
                        operation="fill_empty",
                        match="",
                        replacement="\\sum F_x=0",
                    )
                ]
            )
        )

        result = self.service.repair_study_guide(
            candidate,
            self.sources,
            [
                {
                    "category": "coverage",
                    "location": "sections[0].blocks[1].steps[0].expression",
                    "finding": "Add the source-supported relation used in this step.",
                }
            ],
        )

        self.assertIsNone(candidate["sections"][0]["blocks"][1]["steps"][0]["expression"])
        self.assertEqual(
            result["sections"][0]["blocks"][1]["steps"][0]["expression"],
            "\\sum F_x=0",
        )

    def test_repair_narrows_an_invalid_leaf_to_its_containing_block(self) -> None:
        candidate_section = section(self.upload_id).model_dump(mode="json")
        candidate_section["blocks"].append(
            {
                "type": "example_steps",
                "title": "Short derivation",
                "steps": [
                    {"title": "First", "explanation": "Start from equilibrium."},
                    {"title": "Second", "explanation": "Rearrange the relation."},
                ],
                "sources": [{"upload_id": str(self.upload_id), "page": 2}],
            }
        )
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [candidate_section],
        }
        def repair_one_value(prompt, _schema, **_kwargs):
            target_json = prompt.split(
                "Editable values and exact paths as JSON:\n", 1
            )[1].split("\n\nFindings for these locations as JSON:\n", 1)[0]
            locations = {item["location"] for item in json.loads(target_json)}
            target = "sections[0].blocks[1].steps[0].explanation"
            if target in locations:
                return _StudyGuideRepairEdits(
                    edits=[
                        _StudyGuideRepairEdit(
                            location=target,
                            operation="replace",
                            match="Start from equilibrium.",
                            replacement="Begin with equilibrium.",
                        )
                    ]
                )
            return _StudyGuideRepairEdits(edits=[])

        self.service._generate_json = Mock(side_effect=repair_one_value)

        result = self.service.repair_study_guide(
            candidate,
            self.sources,
            [
                {
                    "category": "accuracy",
                    "location": "sections[0].blocks[1].text",
                    "finding": "Correct the first step against the cited material.",
                }
            ],
        )

        self.assertGreater(self.service._generate_json.call_count, 0)
        self.assertEqual(
            result["sections"][0]["blocks"][1]["steps"][0]["explanation"],
            "Begin with equilibrium.",
        )
        self.assertEqual(
            result["sections"][0]["blocks"][1]["steps"][1]["explanation"],
            "Rearrange the relation.",
        )

    def test_repair_rejects_edits_outside_the_finding_scope(self) -> None:
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [
                section(self.upload_id, text="Keep this claim.").model_dump(mode="json"),
                section(self.upload_id, heading="Other", text="Fix this claim.").model_dump(
                    mode="json"
                ),
            ],
        }
        self.service._generate_json = Mock(
            return_value=_StudyGuideRepairEdits(
                edits=[
                    _StudyGuideRepairEdit(
                        location="sections[0].blocks[0].text",
                        operation="replace",
                        match="Keep this claim.",
                        replacement="Unauthorized change.",
                    )
                ]
            )
        )

        with self.assertRaises(ContentGenerationFailure) as raised:
            self.service.repair_study_guide(
                candidate,
                self.sources,
                [
                    {
                        "category": "accuracy",
                        "location": "sections[1].blocks[0].text",
                        "finding": "Correct this claim.",
                    }
                ],
            )
        self.assertEqual(raised.exception.diagnostic_code, "quality_repair_invalid_patch")

    def test_repair_fails_closed_for_derived_or_missing_locations(self) -> None:
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [section(self.upload_id).model_dump(mode="json")],
        }
        self.service._generate_json = Mock()

        for findings in (
            [
                {
                    "category": "accuracy",
                    "location": "sections[0].body",
                    "finding": "Correct this claim.",
                }
            ],
            [],
        ):
            with self.subTest(findings=findings):
                with self.assertRaises(ContentGenerationFailure) as raised:
                    self.service.repair_study_guide(candidate, self.sources, findings)
                self.assertEqual(
                    raised.exception.diagnostic_code, "quality_repair_invalid_location"
                )

        self.service._generate_json.assert_not_called()

    def test_repair_can_insert_after_one_exact_text_anchor(self) -> None:
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [
                section(self.upload_id, text="The source supports this result.").model_dump(
                    mode="json"
                )
            ],
        }
        self.service._generate_json = Mock(
            return_value=_StudyGuideRepairEdits(
                edits=[
                    _StudyGuideRepairEdit(
                        location="sections[0].blocks[0].text",
                        operation="insert_after",
                        match="The source supports this result.",
                        replacement=" The next step follows by equilibrium.",
                    )
                ]
            )
        )

        result = self.service.repair_study_guide(
            candidate,
            self.sources,
            [
                {
                    "category": "coverage",
                    "location": "sections[0].blocks[0].text",
                    "finding": "Add the missing reasoning.",
                }
            ],
        )
        self.assertEqual(
            result["sections"][0]["blocks"][0]["text"],
            "The source supports this result. The next step follows by equilibrium.",
        )

    def test_repair_context_uses_a_bounded_cited_source_excerpt(self) -> None:
        large_source = {
            **self.sources[0],
            "extracted_text": "Source-backed course text. " * 200,
        }
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [section(self.upload_id, text="Old claim.").model_dump(mode="json")],
        }
        self.service._generate_json = Mock(
            return_value=_StudyGuideRepairEdits(
                edits=[
                    _StudyGuideRepairEdit(
                        location="sections[0].blocks[0].text",
                        operation="replace",
                        match="Old claim.",
                        replacement="Correct claim.",
                    )
                ]
            )
        )

        with patch("app.services.gemini_content.MAX_STUDY_GUIDE_REPAIR_SOURCE_CHARS", 128):
            self.service.repair_study_guide(
                candidate,
                [large_source],
                [
                    {
                        "category": "accuracy",
                        "location": "sections[0].blocks[0].text",
                        "finding": "Correct this claim against its cited source.",
                    }
                ],
            )

        prompt = self.service._generate_json.call_args.args[0]
        source_json = prompt.split("Cited course excerpts as JSON:\n", 1)[1].split(
            "\n\nEditable values and exact paths as JSON:", 1
        )[0]
        excerpts = json.loads(source_json)
        self.assertEqual(len(excerpts), 1)
        self.assertLessEqual(len(excerpts[0]["text"]), 128)
        self.assertTrue(excerpts[0]["truncated"])

    def test_repair_context_selects_only_the_edited_blocks_cited_slide(self) -> None:
        slide_source = {
            **self.sources[0],
            "extracted_text": (
                "[Slide 1]\nUnrelated introductory material.\n"
                "[Slide 2]\nThe cited equilibrium equation is F = ma.\n"
                "[Slide 3]\nUnrelated closing material."
            ),
        }
        section_data = section(self.upload_id).model_dump(mode="json")
        section_data["blocks"].append(
            {
                "type": "formula",
                "expression": "F=0",
                "sources": [{"upload_id": str(self.upload_id), "slide": 2}],
            }
        )
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [section_data],
        }
        self.service._generate_json = Mock(
            return_value=_StudyGuideRepairEdits(
                edits=[
                    _StudyGuideRepairEdit(
                        location="sections[0].blocks[1].expression",
                        operation="replace",
                        match="F=0",
                        replacement="F=ma",
                    )
                ]
            )
        )

        self.service.repair_study_guide(
            candidate,
            [slide_source],
            [
                {
                    "category": "accuracy",
                    "location": "sections[0].blocks[1].expression",
                    "finding": "Use the equation from the cited slide.",
                }
            ],
        )

        prompt = self.service._generate_json.call_args.args[0]
        source_json = prompt.split("Cited course excerpts as JSON:\n", 1)[1].split(
            "\n\nEditable values and exact paths as JSON:", 1
        )[0]
        excerpts = json.loads(source_json)
        self.assertEqual(excerpts[0]["locator"], {"slide": 2})
        self.assertIn("The cited equilibrium equation is F = ma.", excerpts[0]["text"])
        self.assertNotIn("Unrelated introductory material.", excerpts[0]["text"])
        self.assertNotIn("Unrelated closing material.", excerpts[0]["text"])

    def test_repair_can_fill_an_empty_authored_value_without_rebuilding_its_block(self) -> None:
        section_data = section(self.upload_id).model_dump(mode="json")
        section_data["blocks"].append(
            {
                "type": "example_steps",
                "title": "Boundary check",
                "steps": [
                    {"title": "Start", "explanation": "Use the stated boundary."},
                    {
                        "title": "Check",
                        "explanation": "Apply the source-supported equilibrium relation.",
                        "expression": "",
                    },
                ],
                "sources": [{"upload_id": str(self.upload_id)}],
            }
        )
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [section_data],
        }
        self.service._generate_json = Mock(
            return_value=_StudyGuideRepairEdits(
                edits=[
                    _StudyGuideRepairEdit(
                        location="sections[0].blocks[1].steps[1].expression",
                        operation="fill_empty",
                        match="",
                        replacement="\\int_A \\sigma_{ij} n_j\\,dA = F_i",
                    )
                ]
            )
        )

        result = self.service.repair_study_guide(
            candidate,
            self.sources,
            [
                {
                    "category": "coverage",
                    "location": "sections[0].blocks[1].steps[1].expression",
                    "finding": "Add the missing source-supported relation.",
                }
            ],
        )

        self.assertEqual(
            result["sections"][0]["blocks"][1]["steps"][1]["expression"],
            "\\int_A \\sigma_{ij} n_j\\,dA = F_i",
        )
        self.assertEqual(
            result["sections"][0]["blocks"][0]["text"],
            "Traction is a vector.",
        )
        prompt = self.service._generate_json.call_args.args[0]
        self.assertIn("fill_empty only when the supplied value is exactly empty", prompt)
        self.assertIn('"value": ""', prompt)
        self.assertIn(
            "never copy an unrelated equation from the same source into it", prompt
        )

    def test_fill_empty_rejects_a_nonempty_current_value(self) -> None:
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [
                section(self.upload_id, text="Keep this sentence.").model_dump(mode="json")
            ],
        }
        with self.assertRaises(ContentGenerationFailure) as raised:
            self.service._apply_study_guide_edits(
                candidate,
                {"sections[0].blocks[0].text": "Keep this sentence."},
                ["sections[0].blocks[0].text"],
                [
                    _StudyGuideRepairEdit(
                        location="sections[0].blocks[0].text",
                        operation="fill_empty",
                        match="",
                        replacement="Replacement text.",
                    )
                ],
            )
        self.assertEqual(raised.exception.diagnostic_code, "quality_repair_invalid_patch")

    def test_repair_rejects_empty_required_step_title_atomically(self) -> None:
        section_data = section(self.upload_id).model_dump(mode="json")
        section_data["blocks"].append(
            {
                "type": "example_steps",
                "title": "Worked example",
                "steps": [
                    {"title": "Step one", "explanation": "First step."},
                    {"title": "Step two", "explanation": "Second step."},
                ],
                "sources": [{"upload_id": str(self.upload_id)}],
            }
        )
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [section_data],
        }
        original = json.loads(json.dumps(candidate))

        with self.assertRaises(ContentGenerationFailure) as raised:
            self.service._apply_study_guide_edits(
                candidate,
                {"sections[0].blocks[1].steps[0].title": "Step one"},
                ["sections[0].blocks[1].steps[0].title"],
                [
                    _StudyGuideRepairEdit(
                        location="sections[0].blocks[1].steps[0].title",
                        operation="replace",
                        match="Step one",
                        replacement="",
                    )
                ],
            )

        self.assertEqual(raised.exception.diagnostic_code, "quality_repair_invalid_patch")
        self.assertEqual(candidate, original)

    def test_factual_repairs_group_multiple_authored_values_per_request(self) -> None:
        section_data = section(
            self.upload_id, text="A vector is not a tensor."
        ).model_dump(mode="json")
        section_data["blocks"].append(
            {
                "type": "formula",
                "expression": "E=mc^2",
                "sources": [{"upload_id": str(self.upload_id)}],
            }
        )
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [section_data],
        }
        self.service._generate_json = Mock(
            return_value=_StudyGuideRepairEdits(
                edits=[
                    _StudyGuideRepairEdit(
                        location="sections[0].blocks[0].text",
                        operation="replace",
                        match="A vector is not a tensor.",
                        replacement="A vector and a tensor are distinct objects.",
                    ),
                    _StudyGuideRepairEdit(
                        location="sections[0].blocks[1].expression",
                        operation="replace",
                        match="E=mc^2",
                        replacement="E=mc^2+0",
                    ),
                ]
            )
        )

        result = self.service.repair_study_guide(
            candidate,
            self.sources,
            [
                {
                    "category": "accuracy",
                    "location": "sections[0].blocks[0].text",
                    "finding": "Correct the text.",
                },
                {
                    "category": "notation",
                    "location": "sections[0].blocks[1].expression",
                    "finding": "Correct the equation.",
                },
            ],
        )
        self.assertEqual(self.service._generate_json.call_count, 1)
        prompt = self.service._generate_json.call_args.args[0]
        self.assertIn("sections[0].blocks[0].text", prompt)
        self.assertIn("sections[0].blocks[1].expression", prompt)
        for call in self.service._generate_json.call_args_list:
            self.assertEqual(call.kwargs["max_output_tokens"], 6_000)
        self.assertEqual(result["sections"][0]["blocks"][1]["expression"], "E=mc^2+0")

    def test_repair_rejects_ambiguous_text_spans(self) -> None:
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [
                section(
                    self.upload_id,
                    text="Unsupported claim. Repeated phrase. Repeated phrase.",
                ).model_dump(mode="json")
            ],
        }
        self.service._generate_json = Mock(
            return_value=_StudyGuideRepairEdits(
                edits=[
                    _StudyGuideRepairEdit(
                        location="sections[0].blocks[0].text",
                        operation="replace",
                        match="Repeated phrase.",
                        replacement="Correct phrase.",
                    )
                ]
            )
        )

        with self.assertRaises(ContentGenerationFailure) as raised:
            self.service.repair_study_guide(
                candidate,
                self.sources,
                [
                    {
                        "category": "accuracy",
                        "location": "sections[0].blocks[0].text",
                        "finding": "Correct the repeated phrase.",
                    }
                ],
            )
        self.assertEqual(raised.exception.diagnostic_code, "quality_repair_invalid_patch")

    def test_numeric_visual_parameter_can_be_corrected_in_place(self) -> None:
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [
                {
                    **section(self.upload_id).model_dump(mode="json"),
                    "blocks": [
                        section(self.upload_id).model_dump(mode="json")["blocks"][0],
                        {
                            "type": "interactive_model",
                            "template_id": "linear_graph_v1",
                            "title": "Source-supported graph",
                            "parameters": {"slope": 1.0, "intercept": 0.0},
                            "sources": [{"upload_id": str(self.upload_id)}],
                        },
                    ],
                }
            ],
        }
        self.service._generate_json = Mock(
            return_value=_StudyGuideRepairEdits(
                edits=[
                    _StudyGuideRepairEdit(
                        location="sections[0].blocks[1].parameters.slope",
                        operation="replace",
                        match="1.0",
                        replacement="2.0",
                    )
                ]
            )
        )

        result = self.service.repair_study_guide(
            candidate,
            self.sources,
            [
                {
                    "category": "calculation",
                    "location": "sections[0].blocks[1].parameters.slope",
                    "finding": "Use the supported slope.",
                }
            ],
        )
        self.assertEqual(result["sections"][0]["blocks"][1]["parameters"]["slope"], 2.0)

    def test_oversized_guide_repairs_in_small_batches_with_size_limits(self) -> None:
        verbose = "Detailed source-supported explanation. " * 180
        candidate = {
            "format": "study_guide",
            "schema_version": 2,
            "sections": [
                section(
                    self.upload_id,
                    heading=f"Topic {index + 1}",
                    text=verbose,
                ).model_dump(mode="json", exclude_none=True)
                for index in range(20)
            ],
        }
        responses = [
            _StudyGuideRepairEdits(
                edits=[
                    _StudyGuideRepairEdit(
                        location=f"sections[{index}].blocks[0].text",
                        operation="replace",
                        match=verbose.strip(),
                        replacement="Concise source-supported explanation.",
                    )
                ]
            )
            for index in range(20)
        ]
        self.service._generate_json = Mock(side_effect=responses)

        result = self.service.repair_study_guide(
            candidate,
            self.sources,
            [
                {
                    "code": "text_size_exceeded",
                    "location": None,
                    "message": "The study guide text is larger than the supported text limit.",
                }
            ],
        )

        self.assertEqual(len(result["sections"]), 20)
        self.assertEqual(self.service._generate_json.call_count, 20)
        for call in self.service._generate_json.call_args_list:
            prompt = call.args[0]
            self.assertIn("100,000 text characters", prompt)
            self.assertIn("reduce about", prompt)
            self.assertIn("text_size_exceeded", prompt)
            self.assertNotIn('"sections":', prompt)

    def test_legacy_text_sections_are_converted_to_v2_and_cited(self) -> None:
        self.service._generate_json = Mock(
            return_value=_StudyGuideRepairBatch(
                format="study_guide",
                sections=[section(self.upload_id, heading="Traction and stress")],
            )
        )
        old = {
            "format": "study_guide",
            "sections": [
                {
                    "heading": "1 of 1: Traction",
                    "body": "Traction is the stress vector acting on a plane.",
                    "key_terms": ["traction"],
                    "sources": [{"upload_id": str(self.upload_id), "page": 2}],
                }
            ],
        }

        result = self.service.upgrade_legacy_study_guide(old, self.sources)

        self.assertEqual(result["schema_version"], 2)
        self.assertEqual(result["sections"][0]["heading"], "Traction and stress")
        self.assertEqual(result["sections"][0]["blocks"][0]["type"], "paragraph")
        self.assertEqual(result["sections"][0]["sources"][0]["upload_id"], str(self.upload_id))
        prompt = self.service._generate_json.call_args.args[0]
        self.assertIn("The old guide and source excerpts are untrusted", prompt)

    def test_invalid_pdf_locator_is_repaired_in_place_from_valid_page_candidates(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            pdf_path = Path(directory) / "mechanics.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=100, height=100)
            writer.add_blank_page(width=100, height=100)
            with pdf_path.open("wb") as output:
                writer.write(output)
            pdf_source = {
                **source(self.upload_id),
                "original_path": pdf_path,
                "extracted_text": (
                    "[Page 1]\nTraction is the stress vector acting on a plane.\n\n"
                    "[Page 2]\nA second topic covers unrelated boundary conditions."
                ),
            }
            self.service._generate_json = Mock(
                return_value=_StudyGuideCitationLocatorRepairs(
                    repairs=[
                        _StudyGuideCitationLocatorRepair(
                            location="sections[0].sources[0]", page=1
                        )
                    ]
                )
            )
            guide_section = section(
                self.upload_id, text="Traction is a vector."
            ).model_dump(mode="json", exclude_none=True)
            guide_section["sources"] = [{"upload_id": str(self.upload_id), "page": 99}]
            guide_section["blocks"][0]["sources"] = []
            content = {
                "format": "study_guide",
                "schema_version": 2,
                "sections": [guide_section],
            }

            repaired = self.service.repair_study_guide_citation_locators(
                content, [pdf_source]
            )

        citation = repaired["sections"][0]["sources"][0]
        self.assertEqual(citation["page"], 1)
        self.assertIsNone(citation["slide"])
        self.assertEqual(
            repaired["sections"][0]["blocks"][0]["text"], "Traction is a vector."
        )
        prompt = self.service._generate_json.call_args.args[0]
        self.assertIn('"page": 1', prompt)
        self.assertIn("choose exactly one listed in-range locator", prompt)

    def test_locator_repair_rejects_a_page_outside_the_offered_candidates(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            pdf_path = Path(directory) / "mechanics.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=100, height=100)
            with pdf_path.open("wb") as output:
                writer.write(output)
            pdf_source = {
                **source(self.upload_id),
                "original_path": pdf_path,
                "extracted_text": "[Page 1]\nTraction is a stress vector.",
            }
            self.service._generate_json = Mock(
                return_value=_StudyGuideCitationLocatorRepairs(
                    repairs=[
                        _StudyGuideCitationLocatorRepair(
                            location="sections[0].sources[0]", page=999
                        )
                    ]
                )
            )
            guide_section = section(self.upload_id).model_dump(
                mode="json", exclude_none=True
            )
            guide_section["sources"] = [{"upload_id": str(self.upload_id), "page": 9}]
            guide_section["blocks"][0]["sources"] = []
            content = {
                "format": "study_guide",
                "schema_version": 2,
                "sections": [guide_section],
            }

            with self.assertRaises(ContentGenerationFailure) as raised:
                self.service.repair_study_guide_citation_locators(content, [pdf_source])

        self.assertEqual(
            raised.exception.diagnostic_code, "quality_citation_locator_invalid_patch"
        )
        self.assertEqual(content["sections"][0]["sources"][0]["page"], 9)

    def test_locator_repair_keeps_explicit_block_citations_independent(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            pdf_path = Path(directory) / "mechanics.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=100, height=100)
            writer.add_blank_page(width=100, height=100)
            with pdf_path.open("wb") as output:
                writer.write(output)
            pdf_source = {
                **source(self.upload_id),
                "original_path": pdf_path,
                "extracted_text": (
                    "[Page 1]\nTraction is the stress vector acting on a plane.\n\n"
                    "[Page 2]\nBoundary conditions prescribe displacement or traction."
                ),
            }
            self.service._generate_json = Mock(
                side_effect=[
                    _StudyGuideCitationLocatorRepairs(
                        repairs=[
                            _StudyGuideCitationLocatorRepair(
                                location="sections[0].sources[0]", page=1
                            ),
                            _StudyGuideCitationLocatorRepair(
                                location="sections[0].blocks[0].sources[0]", page=1
                            ),
                        ]
                    ),
                    _StudyGuideCitationLocatorRepairs(
                        repairs=[
                            _StudyGuideCitationLocatorRepair(
                                location="sections[0].blocks[1].sources[0]", page=2
                            )
                        ]
                    ),
                ]
            )
            citation = {"upload_id": str(self.upload_id), "page": 99}
            content = {
                "format": "study_guide",
                "schema_version": 2,
                "sections": [
                    {
                        "heading": "Traction and boundaries",
                        "sources": [dict(citation)],
                        "blocks": [
                            {
                                "type": "paragraph",
                                "text": "Traction is the stress vector acting on a plane.",
                                "sources": [dict(citation)],
                            },
                            {
                                "type": "paragraph",
                                "text": "Boundary conditions prescribe displacement or traction.",
                                "sources": [dict(citation)],
                            },
                        ],
                    }
                ],
            }

            repaired = self.service.repair_study_guide_citation_locators(
                content, [pdf_source]
            )

        section_data = repaired["sections"][0]
        self.assertEqual(section_data["sources"][0]["page"], 1)
        self.assertEqual(section_data["blocks"][0]["sources"][0]["page"], 1)
        self.assertEqual(section_data["blocks"][1]["sources"][0]["page"], 2)
        prompt_targets = "\n".join(
            call.args[0].split("Targets and allowed evidence excerpts as JSON:\n", 1)[1]
            for call in self.service._generate_json.call_args_list
        )
        self.assertEqual(prompt_targets.count('"location":'), 3)

    def test_in_range_pdf_locator_does_not_call_model_for_repair(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            pdf_path = Path(directory) / "mechanics.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=100, height=100)
            with pdf_path.open("wb") as output:
                writer.write(output)
            pdf_source = {
                **source(self.upload_id),
                "original_path": pdf_path,
                "extracted_text": "[Page 1]\nTraction is a stress vector.",
            }
            self.service._generate_json = Mock()
            content = {
                "format": "study_guide",
                "schema_version": 2,
                "sections": [
                    {
                        **section(self.upload_id, text="Traction is a vector.").model_dump(
                            mode="json", exclude_none=True
                        ),
                        "sources": [{"upload_id": str(self.upload_id), "page": 1}],
                    }
                ],
            }

            result = self.service.repair_study_guide_citation_locators(
                content, [pdf_source]
            )

        self.assertEqual(result["sections"][0]["sources"][0]["page"], 1)
        self.service._generate_json.assert_not_called()

    def test_pptx_page_locator_is_repaired_to_an_in_range_slide(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            presentation_path = Path(directory) / "mechanics.pptx"
            presentation = Presentation()
            presentation.slides.add_slide(presentation.slide_layouts[6])
            presentation.save(presentation_path)
            slide_source = {
                **source(self.upload_id),
                "original_path": presentation_path,
                "extracted_text": "[Slide 1]\nThe stress diagram shows traction on a plane.",
            }
            self.service._generate_json = Mock(
                return_value=_StudyGuideCitationLocatorRepairs(
                    repairs=[
                        _StudyGuideCitationLocatorRepair(
                            location="sections[0].sources[0]", slide=1
                        )
                    ]
                )
            )
            guide_section = section(
                self.upload_id, text="Traction acts on a plane."
            ).model_dump(mode="json", exclude_none=True)
            guide_section["sources"] = [{"upload_id": str(self.upload_id), "page": 1}]
            guide_section["blocks"][0]["sources"] = []
            content = {
                "format": "study_guide",
                "schema_version": 2,
                "sections": [guide_section],
            }

            repaired = self.service.repair_study_guide_citation_locators(
                content, [slide_source]
            )

        citation = repaired["sections"][0]["sources"][0]
        self.assertIsNone(citation["page"])
        self.assertEqual(citation["slide"], 1)

    def test_visual_render_failure_includes_bounded_locator_diagnostic(self) -> None:
        with patch(
            "app.services.gemini_content.extract_cited_source_visuals",
            side_effect=VisualEvidenceUnavailableError(
                "the cited page could not be rendered",
                source_format=".pdf",
                locator_kind="page",
                locator_index=7,
            ),
        ):
            with self.assertRaises(ContentGenerationFailure) as raised:
                self.service._quality_multimodal_content("prompt", [], self.sources)

        self.assertEqual(
            raised.exception.diagnostic_code, "quality_visual_evidence_unavailable"
        )
        self.assertEqual(
            raised.exception.validation_issues,
            [
                {
                    "location": "cited_visual_evidence",
                    "type": "render_unavailable",
                    "source_format": ".pdf",
                    "locator_kind": "page",
                    "locator_index": 7,
                }
            ],
        )


if __name__ == "__main__":
    unittest.main()
