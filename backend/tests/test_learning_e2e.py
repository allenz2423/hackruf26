"""Integrated source-to-progress regression for the learning system."""

from __future__ import annotations

import json
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import AsyncMock, patch
from uuid import UUID, uuid4

from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches

from app.api.class_library import _reprocess_upload_job
from app.api.schemas import (
    ConceptGraphAIOutput,
    QuestionBankAIGeneration,
    QuestionBankAIItem,
    QuestionBankAIReview,
    QuestionBankAIReviewItem,
)
from app.config import Settings
from app.main import create_app
from app.services.auth import SESSION_COOKIE
from app.services.quiz_reasoning import ReasoningDecision, ReasoningRubricResult

ORIGIN = "http://localhost:5173"
CORRECT_CHOICE = "DNA polymerase extends the new DNA strand."
CORRECT_CHOICE_ZH = "它向新 DNA 链添加核苷酸。"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def biology_lecture_pptx(lecture_number: int, *, revised: bool = False) -> bytes:
    """Build a real slide deck fixture so the flow checks native slide provenance."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = f"DNA replication lecture {lecture_number}"
    textbox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
    detail = " DNA ligase seals fragments." if revised else ""
    textbox.text = (
        "DNA polymerase adds nucleotides to a growing DNA strand. "
        f"This is lecture-specific detail {lecture_number}.{detail}"
    )
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


class DeterministicLearningProvider:
    """Provider fake that grounds graph and question output in supplied source IDs."""

    model_name = "test-learning-provider-v1"

    def __init__(self) -> None:
        self.graph_contexts: list[dict] = []
        self.generation_contexts: list[dict] = []
        self.review_contexts: list[dict] = []

    def process_text(self, text: str, _kind: str) -> dict:
        return {
            "extracted_text": text,
            "summary": "DNA replication source",
            "key_topics": ["DNA replication"],
            "transcript_segments": [],
            "duration_seconds": None,
            "syllabus_outline": None,
        }

    def process_upload(self, path: Path, kind: str, _mime_type: str) -> dict:
        presentation = Presentation(str(path))
        text = "\n".join(
            shape.text.strip()
            for slide in presentation.slides
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and shape.text.strip()
        )
        return self.process_text(text, kind)

    def extract_concept_graph(self, context: dict) -> ConceptGraphAIOutput:
        self.graph_contexts.append(context)
        chunk_id = context["source_chunks"][0]["chunk_revision_id"]
        return ConceptGraphAIOutput.model_validate(
            {
                "concepts": [
                    {
                        "key": "dna-replication",
                        "canonical_name": "DNA replication",
                        "description": "The process that copies DNA before cell division.",
                        "aliases": [],
                        "parent_key": None,
                        "existing_concept_id": None,
                        "provenance": "explicit",
                        "source_chunk_revision_ids": [chunk_id],
                    },
                    {
                        "key": "dna-replication-normalized-duplicate",
                        "canonical_name": "DNA   REPLICATION",
                        "description": "A duplicate normalized concept proposal.",
                        "aliases": [],
                        "parent_key": None,
                        "existing_concept_id": None,
                        "provenance": "explicit",
                        "source_chunk_revision_ids": [chunk_id],
                    },
                ],
                "edges": [],
            }
        )

    def generate_question_bank(self, context: dict) -> QuestionBankAIGeneration:
        self.generation_contexts.append(context)
        questions = []
        for slot in context["assignments"]:
            level = slot["level"]
            source_id = slot["source_chunk_revision_ids"][0]
            language = context["requested_language"]
            if language == "zh-Hans":
                prompt = f"第 {level} 级：DNA 聚合酶的作用是什么？"
                correct_choice = CORRECT_CHOICE_ZH
                distractor = "核糖体负责复制 DNA 染色体。"
            else:
                prompt = f"At demand level {level}, what does DNA polymerase do?"
                correct_choice = CORRECT_CHOICE
                distractor = "Ribosomes copy the DNA chromosomes."
            questions.append(
                QuestionBankAIItem(
                    assignment_id=slot["assignment_id"],
                    prompt=prompt,
                    kind=slot["kind"],
                    choices=[
                        {"id": "a", "text": correct_choice},
                        {"id": "b", "text": distractor},
                    ],
                    answer="a",
                    accepted_answers=[],
                    explanation="DNA polymerase adds nucleotides to a growing DNA strand.",
                    rubric="Identify the role of DNA polymerase.",
                    level=level,
                    language=language,
                    assessment_target=f"DNA polymerase role {slot['assignment_id']}",
                    concept_ids=[slot["concept_id"]],
                    source_chunk_revision_ids=[source_id],
                )
            )
        return QuestionBankAIGeneration(questions=questions)

    def validate_question_bank(self, context: dict) -> QuestionBankAIReview:
        self.review_contexts.append(context)
        return QuestionBankAIReview(
            reviews=[
                QuestionBankAIReviewItem(
                    index=index,
                    source_supported=True,
                    concept_aligned=True,
                    answer_correct=True,
                    distractors_valid=True,
                    unambiguous=True,
                    level_aligned=True,
                    language_valid=True,
                    leakage_free=True,
                    duplicate_family=False,
                    reason="The item is source-grounded and has a distinct validated target.",
                )
                for index, _ in enumerate(context["candidates"])
            ]
        )


class DeterministicLadderReasoningService:
    def __init__(self) -> None:
        self.decisions = [
            ReasoningDecision(
                status="insufficient",
                reason="The student has not described the process yet.",
                follow_up="What does the enzyme do to the new strand?",
            ),
            ReasoningDecision(
                status="sufficient",
                reason="The student made a meaningful conceptual attempt.",
                follow_up=None,
            ),
            ReasoningDecision(
                status="sufficient",
                reason="The student meaningfully explained the recovery question.",
                follow_up=None,
            ),
        ]
        self.contexts: list[dict] = []
        self.rubric_calls = 0

    def evaluate(self, context: dict) -> ReasoningDecision:
        self.contexts.append(context)
        return self.decisions.pop(0)

    def evaluate_rubric(self, _context: dict) -> ReasoningRubricResult:
        self.rubric_calls += 1
        if self.rubric_calls == 1:
            return ReasoningRubricResult(status="failed", feedback="The explanation is flawed.")
        if self.rubric_calls == 2:
            raise RuntimeError("simulated provider outage")
        return ReasoningRubricResult(status="passed", feedback="The reasoning supports the answer.")


class LearningEndToEndTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        settings = Settings(
            _env_file=None,
            gemini_api_key="test-key-not-used",
            file_url_signing_secret="test-signing-secret-with-at-least-32-characters",
            sqlite_path=":memory:",
            storage_dir=self.temp.name,
            signup_mode="open",
            lean_tool_enabled=False,
        )
        self.client = TestClient(create_app(settings_override=settings), headers={"origin": ORIGIN})
        self.client.__enter__()
        self.addCleanup(self.client.__exit__, None, None, None)
        self.db = self.client.app.state.database
        self.provider = DeterministicLearningProvider()
        self.client.app.state.gemini_content_service = self.provider
        self.cookies: dict[str, str] = {}
        self.users: dict[str, dict] = {}
        self.owner = self.add_user("learning_e2e_owner")
        created = self.call("post", "/api/classes", "learning_e2e_owner", json={"name": "DNA"})
        self.assertEqual(created.status_code, 201, created.text)
        self.class_id = created.json()["id"]
        self.student = self.add_user("learning_e2e_student")
        self.observer = self.add_user("learning_e2e_observer")
        self.add_reader("learning_e2e_student")
        self.add_reader("learning_e2e_observer")

    def add_user(self, username: str) -> dict:
        response = self.client.post(
            "/api/auth/signup", json={"username": username, "password": "learning-e2e-pass"}
        )
        self.assertEqual(response.status_code, 201, response.text)
        self.cookies[username] = response.cookies.get(SESSION_COOKIE)
        self.users[username] = response.json()["user"]
        return self.users[username]

    def call(self, method: str, path: str, username: str, **kwargs):
        return getattr(self.client, method)(
            path, cookies={SESSION_COOKIE: self.cookies[username]}, **kwargs
        )

    def add_reader(self, username: str) -> None:
        invite = self.call(
            "post",
            f"/api/classes/{self.class_id}/invitations",
            "learning_e2e_owner",
            json={"capabilities": ["read"], "max_uses": 1, "expires_in_days": 7},
        )
        self.assertEqual(invite.status_code, 201, invite.text)
        joined = self.call(
            "post", "/api/class-invitations/redeem", username, json={"code": invite.json()["code"]}
        )
        self.assertEqual(joined.status_code, 200, joined.text)

    def test_source_to_private_progress_flow_with_pinned_revision(self) -> None:
        # Upload ten real biology slide decks via the public API. TestClient runs the
        # extraction and indexing background task before returning.
        with patch("app.api.class_library.create_default_document", new_callable=AsyncMock):
            upload = self.call(
                "post",
                f"/api/classes/{self.class_id}/uploads",
                "learning_e2e_owner",
                data={"title": "DNA replication lecture 1", "kind": "slides"},
                files={"file": ("biology-lecture-1.pptx", biology_lecture_pptx(1), PPTX_MIME)},
            )
        self.assertEqual(upload.status_code, 202, upload.text)
        upload_id = upload.json()["id"]
        for lecture_number in range(2, 11):
            with patch("app.api.class_library.create_default_document", new_callable=AsyncMock):
                lecture = self.call(
                    "post",
                    f"/api/classes/{self.class_id}/uploads",
                    "learning_e2e_owner",
                    data={"title": f"DNA replication lecture {lecture_number}", "kind": "slides"},
                    files={
                        "file": (
                            f"biology-lecture-{lecture_number}.pptx",
                            biology_lecture_pptx(lecture_number),
                            PPTX_MIME,
                        )
                    },
                )
            self.assertEqual(lecture.status_code, 202, lecture.text)
        indexed = self.client.portal.call(
            self.db.fetchrow,
            "SELECT status,source_index_status FROM course_notes WHERE id=?1",
            upload_id,
        )
        self.assertEqual(indexed, {"status": "ready", "source_index_status": "ready"})
        revision_state = self.client.portal.call(
            self.db.fetchrow,
            "SELECT n.processing_revision,h.processing_revision AS indexed_revision "
            "FROM course_notes n JOIN source_index_heads h ON h.note_id=n.id WHERE n.id=?1",
            upload_id,
        )
        self.assertEqual(
            revision_state["processing_revision"],
            revision_state["indexed_revision"],
            revision_state,
        )
        original_chunk_id = str(
            self.client.portal.call(
                self.db.fetchval,
                "SELECT id FROM source_chunk_revisions WHERE revision_id="
                "(SELECT revision_id FROM source_index_heads WHERE note_id=?1) LIMIT 1",
                upload_id,
            )
        )
        self.assertTrue(original_chunk_id)
        original_locator = self.client.portal.call(
            self.db.fetchval,
            "SELECT locator FROM source_chunk_revisions WHERE id=?1",
            original_chunk_id,
        )
        self.assertEqual(json.loads(original_locator), {"slide": 1})

        extraction = self.call(
            "post",
            f"/api/classes/{self.class_id}/concept-graph/extractions",
            "learning_e2e_owner",
            json={"upload_id": upload_id},
        )
        self.assertEqual(extraction.status_code, 202, extraction.text)
        extraction_id = extraction.json()["operation_id"]
        extraction_status = self.call(
            "get",
            f"/api/classes/{self.class_id}/concept-graph/extractions/{extraction_id}",
            "learning_e2e_owner",
        )
        self.assertEqual(extraction_status.status_code, 200, extraction_status.text)
        self.assertEqual(extraction_status.json()["status"], "complete")
        normalized_concept_count = self.client.portal.call(
            self.db.fetchval,
            "SELECT count(*) FROM class_concepts WHERE class_id=?1 "
            "AND normalized_name='dna replication'",
            self.class_id,
        )
        self.assertEqual(normalized_concept_count, 1)
        review = self.call(
            "get", f"/api/classes/{self.class_id}/concept-graph/review", "learning_e2e_owner"
        )
        self.assertEqual(review.status_code, 200, review.text)
        proposal = next(
            item
            for item in review.json()["concepts"]
            if item["canonical_name"] == "DNA replication"
        )
        self.assertEqual(proposal["state"], "proposed")
        self.assertEqual(proposal["sources"][0]["chunk_revision_id"], original_chunk_id)
        approved = self.call(
            "post",
            f"/api/classes/{self.class_id}/concepts/{proposal['id']}/decision",
            "learning_e2e_owner",
            json={"decision": "approve", "expected_version": proposal["version"]},
        )
        self.assertEqual(approved.status_code, 200, approved.text)
        self.assertEqual(approved.json()["state"], "validated")
        concept_id = approved.json()["id"]
        self.assertEqual(
            self.provider.graph_contexts[0]["source_chunks"][0]["chunk_revision_id"],
            original_chunk_id,
        )

        generated = self.call(
            "post",
            f"/api/classes/{self.class_id}/question-bank/generate",
            "learning_e2e_owner",
            json={
                "operation_id": str(uuid4()),
                "concept_ids": [concept_id],
                "upload_ids": [upload_id],
                "count": 8,
                "language": "zh-Hans",
            },
        )
        self.assertEqual(generated.status_code, 201, generated.text)
        material_id = generated.json()["material_id"]
        self.assertEqual(len(self.provider.generation_contexts), 1)
        self.assertEqual(len(self.provider.review_contexts), 1)
        self.assertEqual(
            self.provider.generation_contexts[0]["source_chunks"][0]["chunk_revision_id"],
            original_chunk_id,
        )
        publication = self.call(
            "patch",
            f"/api/resources/study_materials/{material_id}/publication",
            "learning_e2e_owner",
            json={"state": "published"},
        )
        self.assertEqual(publication.status_code, 200, publication.text)

        run_operation = str(uuid4())
        run_path = f"/api/study-materials/{material_id}/runs"
        run = self.call(
            "post",
            run_path,
            "learning_e2e_student",
            json={"operation_id": run_operation, "mode": "standard"},
        )
        self.assertEqual(run.status_code, 201, run.text)
        run_payload = run.json()
        self.assertEqual(len(run_payload["questions"]), 8)
        self.assertTrue(
            all("DNA 聚合酶" in question["prompt"] for question in run_payload["questions"])
        )
        public_text = json.dumps(run_payload).casefold()
        self.assertNotIn('"answer"', public_text)
        self.assertNotIn('"rubric"', public_text)
        self.assertNotIn('"explanation"', public_text)
        self.assertTrue(all(question["choices"] for question in run_payload["questions"]))
        run_id = run_payload["run_id"]

        # The run records its exact source revision before any later source changes.
        pinned = self.client.portal.call(
            self.db.fetchrow,
            "SELECT private_snapshot FROM quiz_run_questions WHERE run_id=?1 "
            "ORDER BY ordinal LIMIT 1",
            run_id,
        )
        pinned_snapshot = json.loads(pinned["private_snapshot"])
        self.assertIn(original_chunk_id, pinned_snapshot["source_chunk_revision_ids"])

        answers = {
            question["id"]: next(
                choice["id"]
                for choice in question["choices"]
                if choice["text"] == CORRECT_CHOICE_ZH
            )
            for question in run_payload["questions"]
        }
        submitted = self.call(
            "post",
            f"/api/study-materials/{material_id}/attempts",
            "learning_e2e_student",
            json={"run_id": run_id, "operation_id": str(uuid4()), "answers": answers},
        )
        self.assertEqual(submitted.status_code, 200, submitted.text)
        self.assertEqual(submitted.json()["score_percent"], 100)

        own_progress = self.call("get", "/api/learning/progress", "learning_e2e_student")
        self.assertEqual(own_progress.status_code, 200, own_progress.text)
        self.assertGreater(own_progress.json()["total_xp"], 0)
        student_concept = next(
            (item for item in own_progress.json()["concepts"] if item["id"] == concept_id),
            None,
        )
        self.assertIsNotNone(student_concept, own_progress.text)
        self.assertGreater(student_concept["estimate"]["evidence_count"], 0)
        self.assertGreater(own_progress.json()["recent_accuracy"]["graded_count"], 0)

        other_run = self.call("get", f"{run_path}/{run_id}", "learning_e2e_observer")
        self.assertEqual(other_run.status_code, 404, other_run.text)
        other_progress = self.call("get", "/api/learning/progress", "learning_e2e_observer")
        self.assertEqual(other_progress.status_code, 200, other_progress.text)
        self.assertEqual(other_progress.json()["total_xp"], 0)
        observer_concept = next(
            (item for item in other_progress.json()["concepts"] if item["id"] == concept_id),
            None,
        )
        self.assertIsNotNone(observer_concept, other_progress.text)
        self.assertEqual(observer_concept["status"], "unassessed")
        self.assertEqual(observer_concept["estimate"]["evidence_count"], 0)

        preference = self.call("get", "/api/learning/preferences", "learning_e2e_student")
        self.assertEqual(preference.status_code, 200, preference.text)
        self.assertEqual(preference.json(), {"chosen_goal": "balanced"})
        updated_preference = self.call(
            "put",
            "/api/learning/preferences",
            "learning_e2e_student",
            json={"chosen_goal": "exam_prep"},
        )
        self.assertEqual(updated_preference.status_code, 200, updated_preference.text)
        self.assertEqual(updated_preference.json(), {"chosen_goal": "exam_prep"})
        detail = self.call("get", f"/api/learning/concepts/{concept_id}", "learning_e2e_student")
        self.assertEqual(detail.status_code, 200, detail.text)
        self.assertEqual(detail.json()["concept"]["label"], "DNA replication")
        self.assertTrue(detail.json()["recent_evidence"])
        recommendations = self.call("get", "/api/learning/recommendations", "learning_e2e_student")
        self.assertEqual(recommendations.status_code, 200, recommendations.text)
        self.assertEqual(recommendations.json()["chosen_goal"], "exam_prep")
        self.assertTrue(recommendations.json()["items"])
        recommendation = recommendations.json()["items"][0]
        recommendation_id = recommendation["id"]
        self.assertEqual(recommendation["class_id"], self.class_id)
        self.assertEqual(recommendation["href"], f"#/classes/{self.class_id}/study")
        self.client.portal.call(
            self.db.execute,
            "UPDATE learning_recommendations SET state='pending' WHERE id=?1",
            recommendation_id,
        )
        engaged = self.call(
            "patch",
            f"/api/learning/recommendations/{recommendation_id}/state",
            "learning_e2e_student",
            json={"state": "active"},
        )
        self.assertEqual(engaged.status_code, 200, engaged.text)
        self.assertEqual(engaged.json()["state"], "active")
        private_patch = self.call(
            "patch",
            f"/api/learning/recommendations/{recommendation_id}/state",
            "learning_e2e_observer",
            json={"state": "active"},
        )
        self.assertEqual(private_patch.status_code, 404, private_patch.text)

        # Mastery Ladder reuses the generated, validated items while keeping
        # question grading and progression server-owned.
        ladder_response = self.call(
            "post",
            "/api/mastery-ladders",
            "learning_e2e_student",
            json={
                "operation_id": str(uuid4()),
                "class_id": self.class_id,
                "concept_ids": [concept_id],
                "language": "zh-Hans",
                "grading_policy": "objective_only",
            },
        )
        self.assertEqual(ladder_response.status_code, 201, ladder_response.text)
        ladder = ladder_response.json()
        ladder_id = ladder["ladder_id"]
        self.assertEqual(ladder["status"], "active")
        public_ladder = json.dumps(ladder).casefold()
        self.assertNotIn('"answer"', public_ladder)
        self.assertNotIn('"rubric"', public_ladder)
        self.assertNotIn('"explanation"', public_ladder)
        for level in range(1, 5):
            ladder = self.call(
                "get", f"/api/mastery-ladders/{ladder_id}", "learning_e2e_student"
            ).json()
            ladder_concept = ladder["concepts"][0]
            question = ladder_concept["current_question"]
            self.assertIsNotNone(question)
            self.assertIn(f"第 {level} 级", question["prompt"])
            if level == 1:
                wrong_choice = next(
                    item for item in question["choices"] if item["text"] != CORRECT_CHOICE_ZH
                )
                wrong_primary = self.call(
                    "post",
                    f"/api/mastery-ladders/{ladder_id}/concepts/{concept_id}/response",
                    "learning_e2e_student",
                    json={
                        "operation_id": str(uuid4()),
                        "question_id": question["id"],
                        "answer": wrong_choice["id"],
                        "expected_version": ladder["version"],
                    },
                )
                self.assertEqual(wrong_primary.status_code, 200, wrong_primary.text)
                self.assertEqual(wrong_primary.json()["progression"], "recovery_offered")
                ladder = self.call(
                    "get", f"/api/mastery-ladders/{ladder_id}", "learning_e2e_student"
                ).json()
                recovery_question = ladder["concepts"][0]["current_question"]
                self.assertEqual(ladder["concepts"][0]["current_level"], 1)
                self.assertNotEqual(recovery_question["id"], question["id"])
                item_families = self.client.portal.call(
                    self.db.fetch,
                    "SELECT question_id,family_id FROM mastery_ladder_items "
                    "WHERE ladder_id=?1 AND concept_id=?2 AND level=1",
                    ladder_id,
                    concept_id,
                )
                family_by_question = {row["question_id"]: row["family_id"] for row in item_families}
                self.assertNotEqual(
                    family_by_question[question["id"]], family_by_question[recovery_question["id"]]
                )
                question = recovery_question
            correct_choice = next(
                item for item in question["choices"] if item["text"] == CORRECT_CHOICE_ZH
            )
            answer_result = self.call(
                "post",
                f"/api/mastery-ladders/{ladder_id}/concepts/{concept_id}/response",
                "learning_e2e_student",
                json={
                    "operation_id": str(uuid4()),
                    "question_id": question["id"],
                    "answer": correct_choice["id"],
                    "expected_version": ladder["version"],
                },
            )
            self.assertEqual(answer_result.status_code, 200, answer_result.text)
            self.assertEqual(
                answer_result.json()["progression"], "complete" if level == 4 else "advanced"
            )
        final_ladder = self.call(
            "get", f"/api/mastery-ladders/{ladder_id}", "learning_e2e_student"
        ).json()
        self.assertEqual(final_ladder["status"], "completed")
        self.assertEqual(final_ladder["concepts"][0]["state"], "passed")

        failed_ladder_response = self.call(
            "post",
            "/api/mastery-ladders",
            "learning_e2e_student",
            json={
                "operation_id": str(uuid4()),
                "class_id": self.class_id,
                "concept_ids": [concept_id],
                "language": "zh-Hans",
                "grading_policy": "objective_only",
            },
        )
        self.assertEqual(failed_ladder_response.status_code, 201, failed_ladder_response.text)
        failed_ladder_id = failed_ladder_response.json()["ladder_id"]
        for expected_progression in ("recovery_offered", "failed"):
            failed_ladder = self.call(
                "get", f"/api/mastery-ladders/{failed_ladder_id}", "learning_e2e_student"
            ).json()
            active_question = failed_ladder["concepts"][0]["current_question"]
            wrong_choice = next(
                choice
                for choice in active_question["choices"]
                if choice["text"] != CORRECT_CHOICE_ZH
            )
            failed_response = self.call(
                "post",
                f"/api/mastery-ladders/{failed_ladder_id}/concepts/{concept_id}/response",
                "learning_e2e_student",
                json={
                    "operation_id": str(uuid4()),
                    "question_id": active_question["id"],
                    "answer": wrong_choice["id"],
                    "expected_version": failed_ladder["version"],
                },
            )
            self.assertEqual(failed_response.status_code, 200, failed_response.text)
            self.assertEqual(failed_response.json()["progression"], expected_progression)
        terminal_ladder = self.call(
            "get", f"/api/mastery-ladders/{failed_ladder_id}", "learning_e2e_student"
        ).json()
        self.assertEqual(terminal_ladder["concepts"][0]["state"], "failed")

        skipped_ladder_response = self.call(
            "post",
            "/api/mastery-ladders",
            "learning_e2e_student",
            json={
                "operation_id": str(uuid4()),
                "class_id": self.class_id,
                "concept_ids": [concept_id],
                "language": "zh-Hans",
                "grading_policy": "objective_only",
            },
        )
        self.assertEqual(skipped_ladder_response.status_code, 201, skipped_ladder_response.text)
        skipped_ladder_id = skipped_ladder_response.json()["ladder_id"]
        skip_state = self.call(
            "get", f"/api/mastery-ladders/{skipped_ladder_id}", "learning_e2e_student"
        ).json()
        skipped = self.call(
            "post",
            f"/api/mastery-ladders/{skipped_ladder_id}/concepts/{concept_id}/skip",
            "learning_e2e_student",
            json={
                "operation_id": str(uuid4()),
                "expected_version": skip_state["version"],
            },
        )
        self.assertEqual(skipped.status_code, 200, skipped.text)
        self.assertEqual(skipped.json()["progression"], "skipped")
        skipped_state = self.call(
            "get", f"/api/mastery-ladders/{skipped_ladder_id}", "learning_e2e_student"
        ).json()
        self.assertEqual(skipped_state["concepts"][0]["state"], "skipped")

        # Reasoning-required progression uses the sufficiency gate and the
        # independent post-answer rubric. The saved answer survives rubric retry.
        reasoning_service = DeterministicLadderReasoningService()
        self.client.app.state.quiz_reasoning_service = reasoning_service
        reasoning_ladder_response = self.call(
            "post",
            "/api/mastery-ladders",
            "learning_e2e_student",
            json={
                "operation_id": str(uuid4()),
                "class_id": self.class_id,
                "concept_ids": [concept_id],
                "language": "zh-Hans",
                "grading_policy": "reasoning_required",
            },
        )
        self.assertEqual(reasoning_ladder_response.status_code, 201, reasoning_ladder_response.text)
        reasoning_ladder = reasoning_ladder_response.json()
        reasoning_ladder_id = reasoning_ladder["ladder_id"]
        question = reasoning_ladder["concepts"][0]["current_question"]
        reasoning_path = (
            f"/api/mastery-ladders/{reasoning_ladder_id}/concepts/{concept_id}/reasoning"
        )
        initial_reasoning = self.call(
            "post",
            reasoning_path,
            "learning_e2e_student",
            json={
                "operation_id": str(uuid4()),
                "text": "Polymerase is something in the cell.",
                "expected_version": reasoning_ladder["version"],
            },
        )
        self.assertEqual(initial_reasoning.status_code, 200, initial_reasoning.text)
        self.assertEqual(initial_reasoning.json()["status"], "insufficient")
        self.assertEqual(
            initial_reasoning.json()["follow_up"], "What does the enzyme do to the new strand?"
        )
        recovered_ladder = self.call(
            "get",
            f"/api/mastery-ladders/{reasoning_ladder_id}",
            "learning_e2e_student",
        ).json()
        attempt = next(
            item for item in recovered_ladder["attempts"] if item["question_id"] == question["id"]
        )
        self.assertEqual(attempt["reasoning_status"], "insufficient")
        self.assertEqual(
            attempt["reasoning_follow_up"], "What does the enzyme do to the new strand?"
        )
        accepted_reasoning = self.call(
            "post",
            reasoning_path,
            "learning_e2e_student",
            json={
                "operation_id": str(uuid4()),
                "text": "It extends the new DNA strand by adding nucleotides.",
                "expected_version": recovered_ladder["version"],
            },
        )
        self.assertEqual(accepted_reasoning.status_code, 200, accepted_reasoning.text)
        self.assertEqual(accepted_reasoning.json()["status"], "accepted")
        self.assertEqual(len(reasoning_service.contexts[-1]["reasoning_turns"]), 2)

        current_ladder = self.call(
            "get",
            f"/api/mastery-ladders/{reasoning_ladder_id}",
            "learning_e2e_student",
        ).json()
        response_operation = str(uuid4())
        response_path = f"/api/mastery-ladders/{reasoning_ladder_id}/concepts/{concept_id}/response"
        answer = next(
            choice["id"] for choice in question["choices"] if choice["text"] == CORRECT_CHOICE_ZH
        )
        failed_rubric = self.call(
            "post",
            response_path,
            "learning_e2e_student",
            json={
                "operation_id": response_operation,
                "question_id": question["id"],
                "answer": answer,
                "expected_version": current_ladder["version"],
            },
        )
        self.assertEqual(failed_rubric.status_code, 200, failed_rubric.text)
        self.assertEqual(failed_rubric.json()["progression"], "recovery_offered")
        after_flawed_reasoning = self.call(
            "get",
            f"/api/mastery-ladders/{reasoning_ladder_id}",
            "learning_e2e_student",
        ).json()
        self.assertEqual(after_flawed_reasoning["concepts"][0]["current_level"], 1)
        self.assertEqual(after_flawed_reasoning["concepts"][0]["recovery_state"], "available")
        recovery_question = after_flawed_reasoning["concepts"][0]["current_question"]
        self.assertNotEqual(recovery_question["id"], question["id"])
        recovery_reasoning = self.call(
            "post",
            reasoning_path,
            "learning_e2e_student",
            json={
                "operation_id": str(uuid4()),
                "text": "It adds nucleotides to extend the new strand.",
                "expected_version": after_flawed_reasoning["version"],
            },
        )
        self.assertEqual(recovery_reasoning.status_code, 200, recovery_reasoning.text)
        self.assertEqual(recovery_reasoning.json()["status"], "accepted")
        current_recovery_ladder = self.call(
            "get",
            f"/api/mastery-ladders/{reasoning_ladder_id}",
            "learning_e2e_student",
        ).json()
        response_operation = str(uuid4())
        recovery_answer = next(
            choice["id"]
            for choice in recovery_question["choices"]
            if choice["text"] == CORRECT_CHOICE_ZH
        )
        provider_failure = self.call(
            "post",
            response_path,
            "learning_e2e_student",
            json={
                "operation_id": response_operation,
                "question_id": recovery_question["id"],
                "answer": recovery_answer,
                "expected_version": current_recovery_ladder["version"],
            },
        )
        self.assertEqual(provider_failure.status_code, 503, provider_failure.text)
        after_failure = self.call(
            "get",
            f"/api/mastery-ladders/{reasoning_ladder_id}",
            "learning_e2e_student",
        ).json()
        failed_attempt = next(
            item
            for item in after_failure["attempts"]
            if item["question_id"] == recovery_question["id"]
        )
        self.assertEqual(failed_attempt["state"], "completed")
        self.assertEqual(failed_attempt["rubric_status"], "failed")
        self.assertIsNone(failed_attempt["correctness"])
        self.assertIsNone(failed_attempt["xp_awarded"])
        self.assertEqual(after_failure["concepts"][0]["current_level"], 1)

        retried_rubric = self.call(
            "post",
            response_path,
            "learning_e2e_student",
            json={
                "operation_id": response_operation,
                "question_id": recovery_question["id"],
                "answer": recovery_answer,
                "expected_version": after_failure["version"],
            },
        )
        self.assertEqual(retried_rubric.status_code, 200, retried_rubric.text)
        self.assertEqual(retried_rubric.json()["progression"], "advanced")
        self.assertEqual(reasoning_service.rubric_calls, 3)
        reasoning_run_id = self.client.portal.call(
            self.db.fetchval,
            "SELECT run_id FROM mastery_ladder_sessions WHERE id=?1",
            reasoning_ladder_id,
        )
        attempt_row_count = self.client.portal.call(
            self.db.fetchval,
            "SELECT count(*) FROM assessment_outcomes WHERE run_id=?1",
            reasoning_run_id,
        )
        self.assertEqual(attempt_row_count, 2)

        unavailable = self.call(
            "post",
            "/api/mastery-ladders",
            "learning_e2e_student",
            json={
                "operation_id": str(uuid4()),
                "class_id": self.class_id,
                "concept_ids": [concept_id],
                "language": "fr",
                "grading_policy": "objective_only",
            },
        )
        self.assertEqual(unavailable.status_code, 201, unavailable.text)
        self.assertEqual(unavailable.json()["status"], "unavailable")

        # Advancing the source cannot rewrite the assessment's pinned evidence or questions.
        storage_key = self.client.portal.call(
            self.db.fetchval,
            "SELECT storage_key FROM course_notes WHERE id=?1",
            upload_id,
        )
        stored_deck = self.client.app.state.private_storage.path_for(storage_key)
        stored_deck.write_bytes(biology_lecture_pptx(1, revised=True))
        self.client.portal.call(
            self.db.execute,
            "UPDATE course_notes SET status='processing' WHERE id=?1",
            upload_id,
        )
        processing_revision = self.client.portal.call(
            self.db.fetchval,
            "SELECT processing_revision FROM course_notes WHERE id=?1",
            upload_id,
        )
        with patch("app.api.class_library.create_default_document", new_callable=AsyncMock):
            self.client.portal.call(
                _reprocess_upload_job,
                self.client.app,
                UUID(upload_id),
                False,
                processing_revision,
            )
        current_chunk_id = self.client.portal.call(
            self.db.fetchval,
            "SELECT id FROM source_chunk_revisions WHERE revision_id="
            "(SELECT revision_id FROM source_index_heads WHERE note_id=?1) LIMIT 1",
            upload_id,
        )
        self.assertNotEqual(str(current_chunk_id), original_chunk_id)
        current_locator = self.client.portal.call(
            self.db.fetchval,
            "SELECT locator FROM source_chunk_revisions WHERE id=?1",
            current_chunk_id,
        )
        self.assertEqual(json.loads(current_locator), {"slide": 1})
        still_pinned = self.client.portal.call(
            self.db.fetchval,
            "SELECT private_snapshot FROM quiz_run_questions WHERE run_id=?1 "
            "ORDER BY ordinal LIMIT 1",
            run_id,
        )
        self.assertIn(original_chunk_id, json.loads(still_pinned)["source_chunk_revision_ids"])
        resumed = self.call("get", f"{run_path}/{run_id}", "learning_e2e_student")
        self.assertEqual(resumed.status_code, 200, resumed.text)
        self.assertEqual(
            [question["prompt"] for question in resumed.json()["questions"]],
            [question["prompt"] for question in run_payload["questions"]],
        )


if __name__ == "__main__":
    unittest.main()
