import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Colors
    BG_DARK = RGBColor(11, 15, 25)          # Deep slate navy #0B0F19
    CARD_BG = RGBColor(22, 30, 49)          # Card surface #161E31
    CARD_BORDER = RGBColor(39, 51, 78)      # Card border #27334E
    ACCENT_PURPLE = RGBColor(129, 90, 246)  # Indigo/Purple #815AF6
    ACCENT_CYAN = RGBColor(14, 165, 233)    # Gemini Cyan #0EA5E9
    ACCENT_GREEN = RGBColor(16, 185, 129)   # Mint Emerald #10B981
    ACCENT_AMBER = RGBColor(245, 158, 11)   # Warm Gold #F59E0B
    TEXT_LIGHT = RGBColor(248, 250, 252)    # Main title #F8FAFC
    TEXT_MUTED = RGBColor(148, 163, 184)    # Slate subtext #94A3B8
    TEXT_DIM = RGBColor(100, 116, 139)      # Dim grey #64748B

    ROOT_DIR = os.path.dirname(os.path.abspath(__file__))
    VIDEO_PATH = os.path.join(ROOT_DIR, "demo-video", "out", "jot-demo.mp4")
    STILLS_DIR = os.path.join(ROOT_DIR, "demo-video", "out", "stills")

    def add_slide_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, tag_text, title_text, subtitle_text=None, tag_color=ACCENT_CYAN):
        # Category Tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = tag_text.upper()
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = tag_color

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.65))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_LIGHT

        if subtitle_text:
            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(0.4))
            tf_sub = sub_box.text_frame
            tf_sub.word_wrap = True
            tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = subtitle_text
            p_sub.font.size = Pt(14)
            p_sub.font.color.rgb = TEXT_MUTED

    def add_card(slide, left, top, width, height, title, items, border_color=CARD_BORDER, title_color=TEXT_LIGHT, icon_badge=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = Inches(0.28)
        tf.margin_right = Inches(0.28)
        tf.margin_top = Inches(0.25)
        tf.margin_bottom = Inches(0.2)

        p_t = tf.paragraphs[0]
        header_text = f"{icon_badge}  {title}" if icon_badge else title
        p_t.text = header_text
        p_t.font.size = Pt(17)
        p_t.font.bold = True
        p_t.font.color.rgb = title_color
        p_t.space_after = Pt(10)

        for item in items:
            p = tf.add_paragraph()
            p.text = f"•  {item}"
            p.font.size = Pt(12)
            p.font.color.rgb = TEXT_MUTED
            p.space_after = Pt(6)

        return card

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    add_slide_bg(slide1)

    # Accent decorative gradient simulation bar
    bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.3), Inches(1.2), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT_CYAN
    bar.line.fill.background()

    # Track Badges
    badges = [
        ("HACKRU FALL 2026", ACCENT_AMBER),
        ("GOOGLE GEMINI TRACK", ACCENT_CYAN),
        ("EDUCATION TRACK", ACCENT_PURPLE),
    ]
    cur_x = 0.8
    for b_text, b_col in badges:
        b_shape = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cur_x), Inches(1.6), Inches(2.4), Inches(0.38))
        b_shape.fill.solid()
        b_shape.fill.fore_color.rgb = CARD_BG
        b_shape.line.color.rgb = b_col
        b_shape.line.width = Pt(1.2)
        tf = b_shape.text_frame
        tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
        p = tf.paragraphs[0]
        p.text = b_text
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = b_col
        cur_x += 2.6

    # Main Project Title
    t_box = slide1.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(7.5), Inches(1.6))
    tf = t_box.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "Jot (Screen Tutor)"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_LIGHT

    p2 = tf.add_paragraph()
    p2.text = "The Real-Time Multimodal AI Study Buddy"
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = ACCENT_CYAN
    p2.space_before = Pt(6)

    # Tagline description
    tagline_box = slide1.shapes.add_textbox(Inches(0.8), Inches(4.0), Inches(6.8), Inches(1.8))
    tf = tagline_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = (
        "An in-browser AI tutor powered by Gemini Live that sees the screen you choose to share, "
        "talks through your homework in real time, and places non-destructive visual notes right on your work.\n\n"
        "Starts with Socratic hints, checks your steps proactively, and turns your class syllabi, PDFs, "
        "and lectures into interactive quizzes and flashcards."
    )
    p.font.size = Pt(13.5)
    p.font.color.rgb = TEXT_MUTED

    # Hero Still Image on Right
    img_hero = os.path.join(STILLS_DIR, "t15.png")
    if os.path.exists(img_hero):
        hero_pic = slide1.shapes.add_picture(img_hero, Inches(7.8), Inches(1.8), Inches(4.8), Inches(4.8 * (1080/1920)))
        border_rect = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.8), Inches(1.8), Inches(4.8), Inches(4.8 * (1080/1920)))
        border_rect.fill.background()
        border_rect.line.color.rgb = ACCENT_PURPLE
        border_rect.line.width = Pt(2)

    # Footer
    ft = slide1.shapes.add_textbox(Inches(0.8), Inches(6.6), Inches(11.7), Inches(0.4))
    tf_ft = ft.text_frame
    p_ft = tf_ft.paragraphs[0]
    p_ft.text = "Devpost Submission: HackRU 2026  •  Live Demo Included  •  Built with FastAPI & Google Gemini Live"
    p_ft.font.size = Pt(11)
    p_ft.font.color.rgb = TEXT_DIM

    # =========================================================================
    # SLIDE 2: THE PROBLEM (WHY HOMEWORK HELP IS BROKEN)
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_slide_bg(slide2)
    add_header(slide2, "Inspiration & Problem", "Why Homework Help Is Broken Today", "Traditional tools either cheat for the student or can't see the context.")

    add_card(
        slide2, Inches(0.8), Inches(2.0), Inches(3.6), Inches(4.6),
        "1. LLMs are Visually Blind",
        [
            "Standard chat boxes force students to type out multi-step equations or code bugs manually.",
            "Text descriptions lose crucial spatial context (diagrams, graphs, compiler logs, and notebook steps).",
            "Copy-pasting formulas and homework sheets breaks the natural study flow."
        ],
        border_color=RGBColor(239, 68, 68),
        title_color=RGBColor(248, 113, 113),
        icon_badge="[X]"
    )

    add_card(
        slide2, Inches(4.8), Inches(2.0), Inches(3.6), Inches(4.6),
        "2. The 'Answer-Cheat' Trap",
        [
            "Search engines and chatbots vomit the direct solution immediately.",
            "Students develop an illusion of competence without genuine conceptual understanding.",
            "Zero pedagogical restraint: no guided questioning, no hints, and no active recall."
        ],
        border_color=RGBColor(245, 158, 11),
        title_color=RGBColor(251, 191, 36),
        icon_badge="[!]"
    )

    add_card(
        slide2, Inches(8.8), Inches(2.0), Inches(3.7), Inches(4.6),
        "3. Siloed Course Context",
        [
            "Generic AI tools have no clue about your professor's specific syllabus or notation.",
            "Course slides, past exam PDFs, and audio lectures sit idle in Downloads folders.",
            "Students lack an integrated study partner tailored to each specific class."
        ],
        border_color=RGBColor(168, 85, 247),
        title_color=RGBColor(192, 132, 252),
        icon_badge="[?]"
    )

    # =========================================================================
    # SLIDE 3: THE SOLUTION - MEET JOT (SCREEN TUTOR)
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_slide_bg(slide3)
    add_header(slide3, "The Solution", "Meet Jot: Your Real-Time Screen-Sharing Tutor", "A private, conversational peer that sees your screen and points where it counts.")

    add_card(
        slide3, Inches(0.8), Inches(2.0), Inches(5.6), Inches(4.7),
        "Core Capabilities",
        [
            "In-Browser Screen Sharing: Share a tab, window, or entire monitor with zero external desktop plugins.",
            "Natural Voice Dialogue: Ask aloud or type; Jot listens, speaks back naturally, and supports fluid voice interruption (barge-in).",
            "Visual Pointer & Annotations: The model calls add_annotations to highlight equations, box minus signs, or circle code bugs directly over your preview.",
            "Socratic Hinting: Gives pedagogical nudges first so you do the actual thinking and learning.",
            "Proactive 'Check My Steps': Silently follows your math or code and only interrupts when you make a mistake.",
            "Class Knowledge Hub: Ingests syllabi, lecture slides, and recordings to power custom quizzes, flashcards, and study guides."
        ],
        border_color=ACCENT_CYAN,
        title_color=ACCENT_CYAN,
        icon_badge="★"
    )

    # Screenshot side
    img_point = os.path.join(STILLS_DIR, "t25.png")
    if os.path.exists(img_point):
        slide3.shapes.add_picture(img_point, Inches(6.8), Inches(2.0), Inches(5.7), Inches(5.7 * (1080/1920)))
        rect = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(2.0), Inches(5.7), Inches(5.7 * (1080/1920)))
        rect.fill.background()
        rect.line.color.rgb = ACCENT_PURPLE
        rect.line.width = Pt(2)

    caption = slide3.shapes.add_textbox(Inches(6.8), Inches(5.4), Inches(5.7), Inches(1.3))
    tf_c = caption.text_frame
    tf_c.word_wrap = True
    p = tf_c.paragraphs[0]
    p.text = "Visual Grounding in Action: Jot highlights the exact step where an algebraic sign flipped without modifying or touching the user's host OS."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 4: EMBEDDED DEMO VIDEO (THE HIGHLIGHT SLIDE)
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_slide_bg(slide4)
    add_header(slide4, "Live Demonstration", "Full Walkthrough: Jot in Action", "Click the player below to watch the complete 115-second narrated walkthrough.", ACCENT_GREEN)

    # Decorative backing frame behind video
    v_frame = slide4.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.77), Inches(1.87), Inches(7.86), Inches(7.86 * (1080/1920) + 0.06)
    )
    v_frame.fill.solid()
    v_frame.fill.fore_color.rgb = CARD_BG
    v_frame.line.color.rgb = ACCENT_GREEN
    v_frame.line.width = Pt(2)

    # Embed the actual MP4 video directly into the slide!
    poster_path = os.path.join(STILLS_DIR, "t15.png") if os.path.exists(os.path.join(STILLS_DIR, "t15.png")) else None
    if os.path.exists(VIDEO_PATH):
        slide4.shapes.add_movie(
            VIDEO_PATH,
            Inches(0.8), Inches(1.9), Inches(7.8), Inches(7.8 * (1080/1920)),
            poster_frame_image=poster_path,
            mime_type="video/mp4"
        )

    # Demo Chapters Card on the right
    add_card(
        slide4, Inches(9.0), Inches(1.9), Inches(3.5), Inches(4.9),
        "Demo Chapters",
        [
            "0:00 - Introduction & Motivation",
            "0:15 - Seamless In-Browser Screen Share",
            "0:32 - Visual Annotations ('add_annotations')",
            "0:44 - Socratic Hints over Direct Answers",
            "0:55 - Class Hub & Material Uploads",
            "1:07 - AI Quizzes, Flashcards & Study Guides",
            "1:18 - In-Stage PDF Document Tutoring",
            "1:29 - 'Check My Steps' Proactive Coaching",
            "1:40 - Zero-Retention Privacy Architecture"
        ],
        border_color=ACCENT_GREEN,
        title_color=ACCENT_GREEN,
        icon_badge="▶"
    )

    # =========================================================================
    # SLIDE 5: TRACK 1: GOOGLE GEMINI TRACK DEEP DIVE
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_slide_bg(slide5)
    add_header(slide5, "HackRU Track Deep Dive", "Google Gemini Track: Real-Time Multimodal Power", "Harnessing gemini-3.8-live for bidirectional low-latency audio, video, and tool use.", ACCENT_CYAN)

    add_card(
        slide5, Inches(0.8), Inches(2.0), Inches(3.6), Inches(4.7),
        "Bidirectional Live Stream",
        [
            "Model: Powered by Google's latest gemini-3.8-live multimodal streaming API.",
            "Sub-Second Latency: Continuous bidirectional audio and visual frame streaming over WebSockets.",
            "Native Voice Barge-In: Natural conversational audio allows students to interrupt the tutor mid-explanation without delay or feedback loops."
        ],
        border_color=ACCENT_CYAN,
        title_color=ACCENT_CYAN,
        icon_badge="⚡"
    )

    add_card(
        slide5, Inches(4.8), Inches(2.0), Inches(3.6), Inches(4.7),
        "Tool-Calling: add_annotations",
        [
            "Real-Time Visual Actions: The model invokes add_annotations to draw bounding boxes, underlines, and labels.",
            "Normalized Coordinates: Gemini outputs bounding boxes [ymin, xmin, ymax, xmax] normalized from 0.0 to 1.0.",
            "DOM Overlay Translation: Frontend dynamically aligns marks over the responsive video preview corresponding to the exact captured frame."
        ],
        border_color=ACCENT_PURPLE,
        title_color=ACCENT_PURPLE,
        icon_badge="🎯"
    )

    add_card(
        slide5, Inches(8.8), Inches(2.0), Inches(3.7), Inches(4.7),
        "Contextual Multimodality",
        [
            "Adaptive Frame Rate: Samples video frames dynamically on motion/content changes to optimize bandwidth and tokens.",
            "Multimodal Socratic Grounding: Merges live visual frame tokens with spoken audio tokens in a unified attention stream.",
            "Document & Math OCR: Recognizes LaTeX, messy student handwriting, and code terminal outputs effortlessly."
        ],
        border_color=ACCENT_AMBER,
        title_color=ACCENT_AMBER,
        icon_badge="👁"
    )

    # =========================================================================
    # SLIDE 6: TRACK 2: EDUCATION TRACK DEEP DIVE
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_slide_bg(slide6)
    add_header(slide6, "HackRU Track Deep Dive", "Education Track: Pedagogical Design & Active Learning", "Built from the ground up to empower learning rather than replace critical thinking.", ACCENT_PURPLE)

    add_card(
        slide6, Inches(0.8), Inches(2.0), Inches(5.6), Inches(4.7),
        "Pedagogical Principles",
        [
            "Socratic First-Principles: System prompts explicitly enforce progressive hinting. The tutor asks guiding questions ('What is the sign of that term?') rather than solving problems outright.",
            "Proactive 'Check My Steps': A dedicated coaching loop watches math or code steps silently. If a student makes a miscalculation, Jot chimes in with a timely voice hint before they dig into a rabbit hole.",
            "Anti-Cheating Architecture: Quiz and exam APIs strictly separate questions from answer keys until submission, verifying conceptual mastery.",
            "Active Recall & Spaced Repetition: Interactive Flashcard Player featuring mastered vs. review sorting piles, shuffle mode, and instant review cycles."
        ],
        border_color=ACCENT_PURPLE,
        title_color=ACCENT_PURPLE,
        icon_badge="🎓"
    )

    # Right side: Class hub screenshot + details
    img_study = os.path.join(STILLS_DIR, "t55.png")
    if os.path.exists(img_study):
        slide6.shapes.add_picture(img_study, Inches(6.8), Inches(2.0), Inches(5.7), Inches(5.7 * (1080/1920)))
        rect = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(2.0), Inches(5.7), Inches(5.7 * (1080/1920)))
        rect.fill.background()
        rect.line.color.rgb = ACCENT_AMBER
        rect.line.width = Pt(2)

    c_box = slide6.shapes.add_textbox(Inches(6.8), Inches(5.4), Inches(5.7), Inches(1.3))
    tf_cb = c_box.text_frame
    tf_cb.word_wrap = True
    p = tf_cb.paragraphs[0]
    p.text = "Grounded Class Materials: Ingests syllabi, past exams, and lecture slides to synthesize customized study packs with direct source page citations."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 7: SYSTEM ARCHITECTURE & PRIVACY DESIGN
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_slide_bg(slide7)
    add_header(slide7, "Technical Architecture", "End-to-End System & Zero-Retention Privacy", "FastAPI session gateway, provider-neutral WebSocket proxy, and sandboxed SVG overlay.")

    add_card(
        slide7, Inches(0.8), Inches(2.0), Inches(3.6), Inches(4.7),
        "Browser Frontend",
        [
            "Vanilla JS + Vite: Ultra-lean, lightning-fast client with zero bloated dependencies.",
            "WebRTC DisplayMedia: Native browser screen capture (window, tab, or monitor).",
            "Sandboxed SVG Overlays: Annotations render strictly within the page over the preview; no external desktop injection or OS hooking.",
            "Provider-Neutral Client: Browser never sees model names, system prompts, or Gemini API keys."
        ],
        border_color=ACCENT_CYAN,
        title_color=ACCENT_CYAN,
        icon_badge="💻"
    )

    add_card(
        slide7, Inches(4.8), Inches(2.0), Inches(3.6), Inches(4.7),
        "FastAPI Backend",
        [
            "Session Gateway: Issues short-lived, one-time WebSocket connection tickets.",
            "Streaming Proxy: Relays sampled video frames and audio chunks directly into Google GenAI SDK.",
            "SQLite Persistence: Stores user class libraries, syllabi metadata, and generated quizzes locally with zero external DB dependencies.",
            "Tool Execution Engine: Validates and marshals visual annotation payloads."
        ],
        border_color=ACCENT_PURPLE,
        title_color=ACCENT_PURPLE,
        icon_badge="⚙"
    )

    add_card(
        slide7, Inches(8.8), Inches(2.0), Inches(3.7), Inches(4.7),
        "Zero-Retention Privacy",
        [
            "Ephemeral Media: Audio and video frames exist purely in server RAM during live streaming.",
            "Zero Screen Recording: Video feeds and microphone streams are never recorded or written to disk.",
            "User-Owned Data: Course files remain strictly in private local storage.",
            "Memory Reset: Instant 'Reset memory' button flushes conversation history on command."
        ],
        border_color=ACCENT_GREEN,
        title_color=ACCENT_GREEN,
        icon_badge="🔒"
    )

    # =========================================================================
    # SLIDE 8: ENGINEERING CHALLENGES & ACCOMPLISHMENTS
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_slide_bg(slide8)
    add_header(slide8, "Engineering Feats", "Challenges We Overcame & What We're Proud Of", "From cross-origin WebSockets to normalized bounding box rendering.")

    add_card(
        slide8, Inches(0.8), Inches(2.0), Inches(5.6), Inches(4.7),
        "Key Technical Challenges",
        [
            "Coordinate Normalization: Matching Gemini's normalized [0,1] visual coordinates to dynamically scaling browser elements with varying aspect ratios and devicePixelRatio.",
            "Full-Duplex Audio & Interruption: Implementing clean barge-in so the student's voice instantly interrupts model speech playback without creating echo feedback.",
            "Dual-Mode Tutoring Engine: Architecting the transition between active conversational Q&A and silent proactive step-monitoring ('Check my steps').",
            "Zero-Leakage Security: Ensuring Google API keys and credentials are 100% backend-owned and never exposed to the client bundle."
        ],
        border_color=RGBColor(245, 158, 11),
        title_color=RGBColor(251, 191, 36),
        icon_badge="⚡"
    )

    add_card(
        slide8, Inches(6.8), Inches(2.0), Inches(5.7), Inches(4.7),
        "Key Accomplishments",
        [
            "Sub-Second Multimodal Responsiveness: Created a snappy, truly conversational experience that feels like sitting next to a human peer tutor.",
            "Non-Destructive Annotation Engine: High-fidelity visual markup overlay that highlights errors without ever modifying the user's host OS or application.",
            "Complete Study Ecosystem: Automated synthesis of study guides, practice tests, and interactive flashcards with source citations from uploaded course materials.",
            "Automated Video Synthesis: Built an end-to-end headless browser rendering pipeline that records full 1080p demo videos directly from code."
        ],
        border_color=ACCENT_GREEN,
        title_color=ACCENT_GREEN,
        icon_badge="🏆"
    )

    # =========================================================================
    # SLIDE 9: FUTURE ROADMAP & SCALABILITY
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_slide_bg(slide9)
    add_header(slide9, "Future Roadmap", "Where We're Taking Jot Next", "Expanding accessibility, collaborative learning, and classroom integrations.")

    add_card(
        slide9, Inches(0.8), Inches(2.0), Inches(3.6), Inches(4.7),
        "Classroom & LMS Sync",
        [
            "Canvas & Blackboard Integration: Direct one-click import of course syllabi, assignments, and lecture slide decks.",
            "Teacher Analytics Dashboard: Anonymized insight reports showing professors the exact concepts students struggle with most.",
            "Custom Teacher Personas: Allow professors to calibrate tutor strictness and pedagogical style."
        ],
        border_color=ACCENT_CYAN,
        title_color=ACCENT_CYAN,
        icon_badge="🏫"
    )

    add_card(
        slide9, Inches(4.8), Inches(2.0), Inches(3.6), Inches(4.7),
        "Multi-Device & Stylus",
        [
            "Tablet & Stylus Mode: Native iPad and Chromebook whiteboard mode with handwriting recognition and digital ink assistance.",
            "Mobile Companion: Use phone as a live document camera pointing down at physical paper homework sheets.",
            "Real-time Scratchpad: Shared collaborative digital canvas inside the app."
        ],
        border_color=ACCENT_PURPLE,
        title_color=ACCENT_PURPLE,
        icon_badge="📱"
    )

    add_card(
        slide9, Inches(8.8), Inches(2.0), Inches(3.7), Inches(4.7),
        "Collaborative Study Pods",
        [
            "Peer Study Rooms: Multi-user study rooms where multiple classmates share screens and solve problem sets together.",
            "AI Co-Facilitator: Jot acts as an impartial peer facilitator, moderating group discussions and nudging the group toward consensus.",
            "Global Language Lock: Real-time bilingual translation for international study groups."
        ],
        border_color=ACCENT_GREEN,
        title_color=ACCENT_GREEN,
        icon_badge="👥"
    )

    # =========================================================================
    # SLIDE 10: DEVPOST SUMMARY & CONCLUSION
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    add_slide_bg(slide10)
    add_header(slide10, "HackRU Fall 2026", "Thank You! Questions & Discussion", "Jot: Share your screen, ask anything, and actually understand it.", ACCENT_AMBER)

    add_card(
        slide10, Inches(0.8), Inches(2.0), Inches(6.0), Inches(4.7),
        "Devpost Submission Summary",
        [
            "Project Name: Jot (Screen Tutor)",
            "Tracks: Google Gemini Track  •  Education Track",
            "Tech Stack: Python, FastAPI, Google GenAI SDK (gemini-3.8-live), SQLite, WebSockets, Vanilla JS, HTML5, CSS3",
            "Live Video: 115s Narrated Walkthrough embedded on Slide 4 (demo-video/out/jot-demo.mp4)",
            "Privacy Standard: Zero screen/audio recording, fully in-memory live streaming",
            "GitHub Repository: Open-source, production-ready codebase with full unit and integration tests"
        ],
        border_color=ACCENT_AMBER,
        title_color=ACCENT_AMBER,
        icon_badge="📄"
    )

    # Right side: Outro card with launch instructions
    add_card(
        slide10, Inches(7.2), Inches(2.0), Inches(5.3), Inches(4.7),
        "Quick Start Instructions",
        [
            "1. Start Backend Gateway:",
            "   backend\\.venv\\Scripts\\python.exe -m uvicorn screen_tutor_backend.app:app --port 8000",
            "2. Start Frontend Server:",
            "   cd frontend && npm run dev",
            "3. Open Browser:",
            "   Visit http://localhost:5173",
            "4. Replay Presentation Video:",
            "   Open Slide 4 in PowerPoint or run demo-video\\out\\jot-demo.mp4"
        ],
        border_color=ACCENT_GREEN,
        title_color=ACCENT_GREEN,
        icon_badge="🚀"
    )

    output_path = os.path.join(ROOT_DIR, "Jot-ScreenTutor-HackRU-Presentation.pptx")
    prs.save(output_path)
    print(f"Successfully generated PowerPoint presentation at: {output_path}")

if __name__ == "__main__":
    build_deck()
